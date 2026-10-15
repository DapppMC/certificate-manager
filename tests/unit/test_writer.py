import openpyxl
import pytest
from certslice.config import load_config
from certslice.writer import MANIFEST_HEADERS, write_workbooks


@pytest.fixture
def cfg():
    return load_config(None)


def test_write_workbooks(tmp_path, cfg):
    groups = [
        # 1. Accepted group with a low-confidence issue date (needs_review)
        {
            "id": 1,
            "pages": [1],
            "status": "accepted",
            "reason": None,
            "nid": "007123",
            "nama": "Budi Santoso",
            "email": "budi@example.com",
            "unit": "UID Jakarta",
            "candidates": [
                {"nid": "007123", "nama": "Budi Santoso", "unit": "UID Jakarta", "score": 100.0}
            ],
            "fields": {
                "title": {"value": "Teknisi Distribusi", "confidence": "high"},
                "issuer": {"value": "LSP PLN", "confidence": "high"},
                "doc_type": {"value": "Sertifikasi", "confidence": "medium"},
                "number": {"value": "123/2025", "confidence": "high"},
                "issue_date": {"value": None, "confidence": "low"},
                "expiry_date": {"value": "2028-08-16", "confidence": "high"},
            },
            "expiry_source": "printed",
            "renewal": None,
            "file_name": "007123_Budi Santoso_Teknisi Distribusi_0000.pdf",
            "source": "text",
        },
        # 2. Unmatched group with 2 candidates across pages [5, 6]
        {
            "id": 2,
            "pages": [5, 6],
            "status": "unmatched",
            "reason": "homonym",
            "nid": None,
            "nama": None,
            "email": None,
            "unit": None,
            "candidates": [
                {"nid": "0011111", "nama": "Ahmad", "unit": "Unit 1", "score": 95.0},
                {"nid": "0022222", "nama": "Ahmad", "unit": "Unit 2", "score": 95.0},
            ],
            "fields": {
                "title": {"value": "Operator Mesin", "confidence": "medium"},
                "issuer": {"value": "LSP Mesin", "confidence": "medium"},
                "doc_type": {"value": "Sertifikasi", "confidence": "medium"},
                "number": {"value": "456/2025", "confidence": "high"},
                "issue_date": {"value": "2025-01-10", "confidence": "high"},
                "expiry_date": {"value": "2028-01-10", "confidence": "high"},
            },
            "expiry_source": "printed",
            "renewal": None,
            "file_name": None,
            "source": "text",
        },
    ]

    write_workbooks(tmp_path, groups, cfg)

    # 1. Verify manifest_slice.xlsx
    wb_m = openpyxl.load_workbook(tmp_path / "manifest_slice.xlsx")
    ws_m = wb_m["manifest"]
    headers_m = [cell.value for cell in ws_m[1]]
    assert headers_m == MANIFEST_HEADERS

    row2_m = ws_m[2]
    assert row2_m[0].value == 1  # No (integer)
    assert row2_m[1].value == "007123"  # NID (text)
    assert row2_m[1].data_type == "s"
    assert row2_m[7].value in (None, "")  # Tanggal Terbit empty
    assert row2_m[8].value == "16/08/2028"  # Tanggal Expired formatted
    assert ws_m.max_row == 2  # Only 1 accepted group

    # 2. Verify unmatched_report.xlsx
    wb_u = openpyxl.load_workbook(tmp_path / "unmatched_report.xlsx")
    assert "unmatched" in wb_u.sheetnames
    assert "needs_review" in wb_u.sheetnames

    ws_u = wb_u["unmatched"]
    row2_u = ws_u[2]
    assert row2_u[0].value == "5-6"  # Halaman label
    assert row2_u[1].value == "homonym"
    assert row2_u[2].value == "Ahmad"
    assert row2_u[3].value == "0011111"
    assert row2_u[6].value == "Ahmad"
    assert row2_u[7].value == "0022222"
    assert row2_u[10].value in (None, "")  # Candidate 3 empty

    ws_nr = wb_u["needs_review"]
    headers_nr = [c.value for c in ws_nr[1]]
    assert headers_nr == ["Nama File PDF", "NID", "Halaman", "Kolom", "Nilai", "Keyakinan"]
    assert ws_nr.max_row == 2
    row2_nr = ws_nr[2]
    assert row2_nr[0].value == "007123_Budi Santoso_Teknisi Distribusi_0000.pdf"
    assert row2_nr[1].value == "007123"
    assert row2_nr[2].value == "1"
    assert row2_nr[3].value == "Tanggal Terbit"
    assert row2_nr[4].value in (None, "")
    assert row2_nr[5].value == "low"

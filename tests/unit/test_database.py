from pathlib import Path
import openpyxl
import pytest

from certslice.config import load_config
from certslice.database import header_key, load_employees
from certslice.errors import InputError


def test_header_key():
    assert header_key("  Nama.Lengkap-Kerja  ") == "nama_lengkap_kerja"
    assert header_key("NID") == "nid"
    assert header_key(None) == ""


def test_load_employees_basic(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["NID", "Nama Lengkap", "UNITKET"])
    ws.append(["007123", "Budi Santoso", "UID Jakarta"])
    ws.append([8812345, "Dewi Lestari", "UID Jabar"])
    path = tmp_path / "basic.xlsx"
    wb.save(path)

    cfg = load_config(None)
    employees = load_employees(path, cfg)
    assert len(employees) == 2
    assert employees[0].nid == "007123"
    assert employees[0].nama == "Budi Santoso"
    assert employees[0].unit == "UID Jakarta"
    assert employees[0].norm == "budi santoso"
    assert employees[1].nid == "8812345"
    assert employees[1].nama == "Dewi Lestari"


def test_load_employees_header_row_3_and_decoy_sheet(tmp_path):
    wb = openpyxl.Workbook()
    # Decoy sheet
    ws_decoy = wb.active
    ws_decoy.title = "Decoy"
    ws_decoy.append(["NID", "Nama"])
    ws_decoy.append(["999", "Decoy Person"])

    # Target sheet
    ws_master = wb.create_sheet(title="Data Master")
    ws_master.append(["Title Row"])
    ws_master.append([])
    ws_master.append(["No", "NID", "Nama Pegawai", "Email Korporat"])
    ws_master.append([1, "00123", "Siti Rahmawati", "siti@contoh.com"])
    path = tmp_path / "decoy.xlsx"
    wb.save(path)

    cfg = load_config(None)
    employees = load_employees(path, cfg)
    assert len(employees) == 1
    assert employees[0].nid == "00123"
    assert employees[0].nama == "Siti Rahmawati"
    assert employees[0].email == "siti@contoh.com"


def test_missing_nid_column(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Nama", "Unit"])
    ws.append(["Budi", "UID"])
    path = tmp_path / "no_nid.xlsx"
    wb.save(path)

    cfg = load_config(None)
    with pytest.raises(InputError) as exc_info:
        load_employees(path, cfg)
    assert "nid" in str(exc_info.value).lower()


def test_empty_rows_below_header(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["NID", "Nama"])
    path = tmp_path / "header_only.xlsx"
    wb.save(path)

    cfg = load_config(None)
    with pytest.raises(InputError) as exc_info:
        load_employees(path, cfg)
    assert "no employee rows" in str(exc_info.value).lower()


def test_corrupt_or_not_xlsx(tmp_path):
    path = tmp_path / "not_an_excel.xlsx"
    path.write_text("random binary garbage", encoding="utf-8")

    cfg = load_config(None)
    with pytest.raises(InputError) as exc_info:
        load_employees(path, cfg)
    assert "cannot open database file" in str(exc_info.value).lower()

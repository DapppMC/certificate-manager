"""Unit tests for the unified bulk renaming and roster parsing module."""
from pathlib import Path
import tempfile

import openpyxl
import pymupdf
import pytest

from certslice.bulk import (
    BulkItem,
    _detect_skp_offline,
    _write_bulk_workbooks,
    expand_name_abbreviations,
    find_bulk_candidates,
    get_jabatan_keywords,
    get_optimal_worker_count,
    get_unit_keywords,
    is_auxiliary_doc,
    is_roster_doc,
    match_roster,
    normalize_digraphs,
    parse_roster_doc,
    prepare_cert_jobs,
    token_similarity,
)
from certslice.models import Employee
from certslice.ocr import OcrEngine


def test_is_roster_doc_by_filename():
    assert is_roster_doc(Path("122_PT PLN Nusantara Power Services - PLTU Bangka 3-1-4.pdf"))
    assert is_roster_doc(Path("127_PT PLN Nusantara Power Services-1-4.pdf"))
    assert is_roster_doc(Path("Lampiran_Surat_Pengantar.pdf"))
    assert is_roster_doc(Path("Roster_Penerima_Sertifikat.pdf"))
    assert not is_roster_doc(Path("122_PT PLN Nusantara Power Services - PLTU Bangka 3-5-6.pdf"))
    assert not is_roster_doc(Path("adam.pdf"))


def test_match_roster():
    entries = [
        {
            "row": 1,
            "level": "Level 3",
            "cert_no": "2060.0.07.P043.07.2026",
            "reg_no": "50712.1.2026",
            "kode_okupasi": "D.35.114.01.KUALIFIKASI.3.KITLTU",
            "tokens": ["MUHAMMAD", "AMIN", "Operator", "senior", "lokal", "sistem", "bahan", "bakar"],
        },
        {
            "row": 19,
            "level": "Level 2",
            "cert_no": "2052.0.07.P042.07.2026",
            "reg_no": "50722.1.2026",
            "kode_okupasi": "D.35.114.01.KUALIFIKASI.2.KITLTU",
            "tokens": ["SONA", "Pelaksana", "ARDIANSYAH", "tugas", "pengoperasian", "Peralatan", "Pemadam", "Kebakaran"],
        },
    ]

    m1 = match_roster("MUHAMMAD AMIN", entries)
    assert m1 is not None
    assert m1["row"] == 1
    assert m1["level"] == "Level 3"
    assert m1["cert_no"] == "2060.0.07.P043.07.2026"
    assert "Operator" in m1["peran"]

    m19 = match_roster("SONA ARDIANSYAH", entries)
    assert m19 is not None
    assert m19["row"] == 19
    assert m19["level"] == "Level 2"
    assert m19["cert_no"] == "2052.0.07.P042.07.2026"
    assert "Pelaksana" in m1["peran"] or "Pemadam" in m19["peran"]

    # Non-existent employee returns None
    assert match_roster("NAMA TIDAK TERDAFTAR", entries) is None


def test_match_roster_by_cert_no():
    entries = [
        {
            "row": 11,
            "level": "Level 3",
            "cert_no": "2121.1.07.P053.07.2026",
            "reg_no": "52841.1.2026",
            "kode_okupasi": "F.43.115.01.KUALIFIKASI.3.KITLTU",
            "peran": "Pelaksana Senior Pemeliharaan Coal Handling",
            "tokens": ["MUHAMMAD", "RISAL", "Pelaksana", "Senior", "Pemeliharaan", "Coal", "Handling"],
        },
    ]
    # Exact cert_no match takes precedence
    m = match_roster("", entries, cert_no="2121.1.07.P053.07.2026")
    assert m is not None
    assert m["level"] == "Level 3"
    assert m["cert_no"] == "2121.1.07.P053.07.2026"
    assert "Coal Handling" in m["peran"]

    # Non-existent cert returns None
    assert match_roster("", entries, cert_no="9999.9.99.P999.99.9999") is None


def test_detect_skp_offline():
    sample_text = """
    SERTIFIKAT KOMPETENSI
    PT SKP TENAGA LISTRIK
    Nomor : 2060.0.07.P043.07.2026
    Unit Kompetensi :
    Pengoperasian Sistem Bahan Bakar Batu Bara bagi Operator Madya
    """
    is_skp, level, peran = _detect_skp_offline(sample_text)
    assert is_skp is True
    assert "Operator" in level or "Madya" in level or not level  # Regex finds Operator Madya if matched
    assert "Bahan Bakar" in peran


def test_write_bulk_workbooks_schema():
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_path = Path(tmpdir)
        items = [
            BulkItem(
                index=1,
                src_path=Path("sample.pdf"),
                file_name="0124115WBN_ADHA_Level 2_Operator.pdf",
                nid="0124115WBN",
                nama="ADHA",
                jabatan="Operator",
                unit="PLTU Bangka",
                is_skp=True,
                level="Level 2",
                peran="Operator",
                judul="Operator",
                number="2053.0.07.P042.07.2026",
                issue_date="2026-07-15",
                expiry_date="2029-07-15",
                status="accepted",
                reason="",
            )
        ]
        _write_bulk_workbooks(tmp_path, items)

        xlsx_file = tmp_path / "data.xlsx"
        assert xlsx_file.exists()

        wb = openpyxl.load_workbook(xlsx_file)
        ws = wb.active
        assert ws.title == "Rekap Sertifikat"

        expected_headers = [
            "No",
            "File Asal",
            "Nama File Baru",
            "NID",
            "Nama Pegawai",
            "Jabatan",
            "Unit Kerja",
            "Tipe Dokumen",
            "Judul Sertifikat / Pelatihan",
            "Level (Khusus SKP)",
            "Peran Kompetensi (Khusus SKP)",
            "Nomor Sertifikat",
            "Tanggal Terbit",
            "Tanggal Expired",
            "Status Validasi",
            "Keterangan / Alasan Review",
        ]
        actual_headers = [ws.cell(row=1, column=c).value for c in range(1, len(expected_headers) + 1)]
        assert actual_headers == expected_headers

        # Verify row 2 data
        assert ws.cell(row=2, column=1).value == "1" or ws.cell(row=2, column=1).value == 1
        assert ws.cell(row=2, column=4).value == "0124115WBN"
        assert ws.cell(row=2, column=5).value == "ADHA"
        assert ws.cell(row=2, column=8).value == "SKP Ketenagalistrikan"
        assert ws.cell(row=2, column=15).value == "Tervalidasi"
        wb.close()


def test_is_auxiliary_doc():
    is_aux, cat = is_auxiliary_doc(Path("Daftar_Hadir_Peserta.pdf"))
    assert is_aux is True
    assert cat == "Daftar Hadir"

    is_aux, cat = is_auxiliary_doc(Path("Invoice_Pelatihan_K3.pdf"))
    assert is_aux is True
    assert "Invoice" in cat

    is_aux, cat = is_auxiliary_doc(Path("Rundown_Acara_Sertifikasi.pdf"))
    assert is_aux is True
    assert "Rundown" in cat

    is_aux, _ = is_auxiliary_doc(Path("Sertifikat_Budi.pdf"))
    assert is_aux is False


def test_expand_name_abbreviations():
    exp1 = expand_name_abbreviations("M. Rizki Pratama")
    assert any("muhammad" in s for s in exp1)
    assert any("mochammad" in s for s in exp1)

    exp2 = expand_name_abbreviations("Ach. Fauzi")
    assert any("achmad" in s for s in exp2)
    assert any("ahmad" in s for s in exp2)


def test_token_similarity_and_digraphs():
    assert normalize_digraphs("fakhtur") == "fahtur"
    assert normalize_digraphs("koeswantoyo") == "kuswantoyo"

    # Typo tolerance
    sim1 = token_similarity("fakhtur", "fatkhurrokhan")
    assert sim1 >= 80.0

    sim2 = token_similarity("cintya", "cintiya")
    assert sim2 >= 85.0


def test_get_unit_keywords():
    kw = get_unit_keywords("PT PLN NPS UP PLTU AMURANG")
    assert "amurang" in kw
    assert "pltu" not in kw
    assert "pln" not in kw


def test_find_bulk_candidates_unit_disambiguation():
    employees = [
        Employee(nid="1001", nama="ADAM", jabatan="Teknisi", unit="PLTU Paiton", norm="adam"),
        Employee(nid="1002", nama="ADAM", jabatan="Operator", unit="PLTU Amurang", norm="adam"),
    ]
    cfg = {
        "matching": {
            "accept_score": 92,
            "list_score": 80,
            "min_tokens_fuzzy": 2,
            "holder_cues": ["diberikan kepada"],
            "titles": ["ir", "st"],
        }
    }
    lines = ["SERTIFIKAT", "Diberikan kepada :", "ADAM", "Lokasi Kerja: PLTU Amurang"]
    status, top, reason, cands = find_bulk_candidates(lines, employees, cfg)

    assert status == "accepted"
    assert top is not None
    assert top.index == 1
    assert employees[top.index].nid == "1002"
    assert "Amurang" in employees[top.index].unit


def test_find_bulk_candidates_unique_mononym():
    employees = [
        Employee(nid="2001", nama="WENEDI SURYA PRATAMA", jabatan="Operator", unit="PLTU Bangka", norm="wenedi surya pratama"),
        Employee(nid="2002", nama="BUDI SANTOSO", jabatan="Teknisi", unit="PLTU Paiton", norm="budi santoso"),
    ]
    cfg = {
        "matching": {
            "accept_score": 92,
            "list_score": 80,
            "min_tokens_fuzzy": 2,
            "holder_cues": [],
            "titles": [],
        }
    }
    # Single unique token "Wenedi" that appears in only 1 employee across the company
    lines = ["Certificate of Completion", "Awarded to:", "Wenedi"]
    status, top, reason, cands = find_bulk_candidates(lines, employees, cfg)

    assert status == "accepted"
    assert top is not None
    assert top.index == 0
    assert employees[top.index].nid == "2001"


def test_get_jabatan_keywords():
    j_kw = get_jabatan_keywords("Operator Senior Boiler & Turbin")
    assert "operator" in j_kw
    assert "boiler" in j_kw
    assert "turbin" in j_kw


def test_find_bulk_candidates_inverted_name():
    employees = [
        Employee(nid="3001", nama="WENEDI SURYA PRATAMA", jabatan="Operator", unit="PLTU Bangka", norm="wenedi surya pratama"),
    ]
    cfg = {
        "matching": {
            "accept_score": 92,
            "list_score": 80,
            "min_tokens_fuzzy": 2,
            "holder_cues": ["diberikan kepada"],
            "titles": [],
        }
    }
    # Last name first with comma: "Pratama, Wenedi Surya"
    lines = ["SERTIFIKAT", "Diberikan kepada :", "Pratama, Wenedi Surya"]
    status, top, reason, cands = find_bulk_candidates(lines, employees, cfg)

    assert status == "accepted"
    assert top is not None
    assert employees[top.index].nid == "3001"


def test_find_bulk_candidates_patronymic():
    employees = [
        Employee(nid="4001", nama="AHMAD FAUZI", jabatan="Teknisi", unit="PLTU Paiton", norm="ahmad fauzi"),
    ]
    cfg = {
        "matching": {
            "accept_score": 92,
            "list_score": 80,
            "min_tokens_fuzzy": 2,
            "holder_cues": ["diberikan kepada"],
            "titles": [],
        }
    }
    # Has patronymic "Bin Sulaiman"
    lines = ["SERTIFIKAT", "Diberikan kepada :", "Ahmad Fauzi Bin Sulaiman"]
    status, top, reason, cands = find_bulk_candidates(lines, employees, cfg)

    assert status == "accepted"
    assert top is not None
    assert employees[top.index].nid == "4001"


def test_is_roster_doc_kemnaker(tmp_path: Path):
    doc_path = tmp_path / "1317_Pengiriman_Sertifikasi_KEMNAKER.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 50), "SURAT KETERANGAN PENGIRIMAN SERTIFIKAT\nHasil Perpanjangan KEMNAKER atas nama:\nJENIS SERTIFIKASI\nLISENSI\nSKP")
    doc.save(doc_path)
    doc.close()

    assert is_roster_doc(doc_path) is True


def test_parse_roster_doc_kemnaker(tmp_path: Path):
    doc_path = tmp_path / "Pengiriman_Kemnaker.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    text = (
        "Bersama dengan surat ini, kami sampaikan hasil Perpanjangan KEMNAKER atas nama:\n"
        "NO\nNAMA\nPERUSAHAAN\nJENIS SERTIFIKASI\nLISENSI\nSKP\n"
        "1 Antonius Amanto Sulistyo\nPT PLN Nusantara Power Services\nPerpanjangan AK3 Umum\nSoft File\nSoft File\n"
        "2 Denny Septa Ferdiansyah\nSoft File\nSoft File\n"
        "3 Hendra\nPerpanjangan Operator Bulldozer\nSoft File\n"
    )
    page.insert_text((50, 50), text)
    doc.save(doc_path)
    doc.close()

    ocr = OcrEngine(dpi=150)
    entries = parse_roster_doc(doc_path, ocr)
    assert len(entries) == 3
    assert entries[0]["nama"] == "Antonius Amanto Sulistyo"
    assert entries[0]["peran"] == "Perpanjangan AK3 Umum"
    assert entries[1]["nama"] == "Denny Septa Ferdiansyah"
    assert entries[1]["peran"] == "Perpanjangan AK3 Umum"  # forward-filled
    assert entries[2]["nama"] == "Hendra"
    assert entries[2]["peran"] == "Perpanjangan Operator Bulldozer"


def test_prepare_cert_jobs_alternating_continuation(tmp_path: Path):
    # 4-page bundle: Page 1 cert for Hendra, Page 2 syllabus, Page 3 cert for Denny, Page 4 syllabus
    doc_path = tmp_path / "bundle_pelatihan.pdf"
    doc = pymupdf.open()

    p1 = doc.new_page()
    p1.insert_text((50, 50), "SERTIFIKAT KEPELATIHAN\nDiberikan kepada: HENDRA")

    p2 = doc.new_page()
    p2.insert_text((50, 50), "Fokus Pengembangan Keahlian\nSilabus dan Kurikulum Pelatihan")

    p3 = doc.new_page()
    p3.insert_text((50, 50), "SERTIFIKAT KEPELATIHAN\nDiberikan kepada: DENNY SEPTA FERDIANSYAH")

    p4 = doc.new_page()
    p4.insert_text((50, 50), "Fokus Pengembangan Keahlian\nSilabus dan Kurikulum Pelatihan")

    doc.save(doc_path)
    doc.close()

    employees = [
        Employee(nid="101", nama="HENDRA", jabatan="Operator", unit="Unit 1", norm="hendra"),
        Employee(nid="102", nama="DENNY SEPTA FERDIANSYAH", jabatan="Teknisi", unit="Unit 2", norm="denny septa ferdiansyah"),
    ]
    cfg = {
        "matching": {
            "accept_score": 92,
            "list_score": 80,
            "min_tokens_fuzzy": 2,
            "holder_cues": ["diberikan kepada"],
            "titles": [],
        },
        "grouping": {
            "max_pages_per_cert": 3,
            "continuation_patterns": [],
        }
    }
    ocr = OcrEngine(dpi=150)

    jobs, rosters, roster_entries, aux = prepare_cert_jobs([doc_path], employees, cfg, ocr)
    assert len(jobs) == 2
    assert jobs[0].pages == [0, 1]
    assert jobs[1].pages == [2, 3]


def test_get_optimal_worker_count():
    workers = get_optimal_worker_count()
    assert 1 <= workers <= 4




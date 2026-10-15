"""Shared helpers for the planner-owned acceptance tests (protected: the executor must not edit
anything under tests/acceptance/ or tests/fixtures/). Contract: docs/spec/modules/interface.md."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))
FIX = ROOT / "tests" / "fixtures"
DB = FIX / "db_basic.xlsx"


def _ensure_fixtures() -> None:
    if (FIX / "bundle_simple.pdf").exists() and (FIX / "db_basic.xlsx").exists():
        return
    sys.path.insert(0, str(ROOT / "scripts"))
    import make_fixtures  # noqa: PLC0415

    make_fixtures.main(["--out", str(FIX)])


_ensure_fixtures()


def cli(argv: list[str], capsys) -> tuple[int, str, str]:
    """Run `python -m certslice <argv>` in-process; return (exit code, stdout, stderr)."""
    from certslice.cli import main  # noqa: PLC0415

    capsys.readouterr()
    try:
        code = main(argv)
    except SystemExit as e:  # argparse errors
        code = e.code if isinstance(e.code, int) else 1
    out, err = capsys.readouterr()
    return code, out, err


def run_dir_from(stdout: str) -> Path:
    m = re.search(r"^RUN_DIR=(.+)$", stdout, re.MULTILINE)
    assert m, f"stdout has no RUN_DIR= line:\n{stdout}"
    return Path(m.group(1).strip())


@pytest.fixture(scope="session")
def outbase(tmp_path_factory) -> Path:
    return tmp_path_factory.mktemp("runs")


_CACHE: dict[tuple, tuple[int, str, Path | None]] = {}


@pytest.fixture
def run_bundle(outbase, capsys):
    """run_bundle("bundle_simple", "--no-ocr") -> (exit code, stdout, run dir). Cached per argument list."""

    def _run(bundle: str, *extra: str, db: Path = DB, cache: bool = True):
        key = (bundle, extra, str(db))
        if cache and key in _CACHE:
            return _CACHE[key]
        out = outbase / re.sub(r"[^A-Za-z0-9]+", "_", "_".join((bundle,) + extra))
        code, so, _ = cli(["run", "--pdf", str(FIX / f"{bundle}.pdf"), "--db", str(db), "--out", str(out), *extra],
                          capsys)
        rd = run_dir_from(so) if code == 0 else None
        res = (code, so, rd)
        if cache:
            _CACHE[key] = res
        return res

    return _run


def load_plan(run_dir: Path) -> dict:
    return json.loads((run_dir / "plan.json").read_text(encoding="utf-8"))


def sheet_rows(path: Path, sheet: str | None = None) -> list[dict]:
    from openpyxl import load_workbook  # noqa: PLC0415

    wb = load_workbook(path)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = list(rows[0])
    return [dict(zip(hdr, r)) for r in rows[1:]]


def manifest(run_dir: Path) -> list[dict]:
    return sheet_rows(run_dir / "manifest_slice.xlsx")


def groups_pages(plan: dict) -> list[list[int]]:
    return [g["pages"] for g in plan["groups"]]


def assert_accounting(run_dir: Path) -> None:
    """R-13: every page in exactly one group; accepted groups have a PDF with that many pages,
    unmatched groups have exactly one report row with the same page string."""
    import pymupdf  # noqa: PLC0415

    plan = load_plan(run_dir)
    seen = [p for g in plan["groups"] for p in g["pages"]]
    assert sorted(seen) == list(range(1, plan["page_count"] + 1)), seen
    unmatched_rows = sheet_rows(run_dir / "unmatched_report.xlsx", "unmatched")
    report_pages = [str(r["Halaman"]) for r in unmatched_rows]
    for g in plan["groups"]:
        pg = g["pages"]
        label = str(pg[0]) if len(pg) == 1 else f"{pg[0]}-{pg[-1]}"
        if g["status"] == "accepted":
            f = run_dir / "pdf" / g["file_name"]
            assert f.exists(), f
            with pymupdf.open(f) as d:
                assert d.page_count == len(pg)
            assert label not in report_pages
        else:
            assert g["status"] == "unmatched"
            assert report_pages.count(label) == 1, (label, report_pages)
    assert len(unmatched_rows) == sum(1 for g in plan["groups"] if g["status"] == "unmatched")

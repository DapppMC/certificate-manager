"""T-014 manifest and report workbooks (R-11, R-12, R-15)."""
from conftest import assert_accounting, manifest, sheet_rows

HEADERS = ["No", "NID", "Nama Pegawai", "Judul Sertifikasi/Pelatihan", "Jenis Dokumen (Sertifikasi/Pelatihan)",
           "Lembaga/Penyelenggara", "Nomor Sertifikat", "Tanggal Terbit", "Tanggal Expired", "Nama File PDF"]


def test_manifest_columns_and_values(run_bundle):
    from openpyxl import load_workbook
    code, out, rd = run_bundle("bundle_simple", "--no-ocr")
    ws = load_workbook(rd / "manifest_slice.xlsx").worksheets[0]
    assert [c.value for c in ws[1]] == HEADERS
    rows = manifest(rd)
    assert len(rows) == 3
    r = rows[0]
    assert r["No"] == 1 and r["NID"] == "007123" and r["Nama Pegawai"] == "Budi Santoso"
    assert r["Judul Sertifikasi/Pelatihan"] == "Teknisi Distribusi Tenaga Listrik"
    assert r["Jenis Dokumen (Sertifikasi/Pelatihan)"] == "Sertifikasi"
    assert r["Lembaga/Penyelenggara"] == "LEMBAGA SERTIFIKASI PROFESI PLN"
    assert r["Nomor Sertifikat"] == "123/LSP-PLN/VIII/2025"
    assert r["Tanggal Terbit"] == "17/08/2025" and r["Tanggal Expired"] == "16/08/2028"
    assert r["Nama File PDF"] == "007123_Budi Santoso_Teknisi Distribusi Tenaga Listrik_2025.pdf"
    for row in ws.iter_rows(min_row=2):
        assert row[1].data_type == "s"  # NID stored as text


def test_unmatched_report(run_bundle):
    code, out, rd = run_bundle("bundle_unmatched", "--no-ocr")
    rows = sheet_rows(rd / "unmatched_report.xlsx", "unmatched")
    by_page = {str(r["Halaman"]): r for r in rows}
    assert set(by_page) == {"1", "2", "4"}
    assert by_page["1"]["Alasan"] == "homonym"
    assert {by_page["1"]["Kandidat 1 NID"], by_page["1"]["Kandidat 2 NID"]} == {"0011111", "0022222"}
    assert by_page["2"]["Alasan"] == "no confident holder"
    assert by_page["4"]["Alasan"] == "close candidates"
    assert len(manifest(rd)) == 1
    assert_accounting(rd)


def test_needs_review_sheet(run_bundle):
    code, out, rd = run_bundle("bundle_mixed", "--no-ocr")
    rows = sheet_rows(rd / "unmatched_report.xlsx", "needs_review")
    hits = [r for r in rows if r["NID"] == "0044444" and r["Kolom"] == "Tanggal Terbit"]
    assert len(hits) == 1
    m = [r for r in manifest(rd) if r["NID"] == "0044444"][0]
    assert m["Tanggal Terbit"] in (None, "")
    assert m["Nama File PDF"].endswith("_0000.pdf")
    assert_accounting(rd)


def test_training_expiry_cell_empty(run_bundle):
    code, out, rd = run_bundle("bundle_internal", "--no-ocr")
    for r in manifest(rd):
        assert r["Jenis Dokumen (Sertifikasi/Pelatihan)"] == "Pelatihan"
        assert r["Tanggal Expired"] in (None, "")


def test_partial_name_row_lists_the_likely_employee(run_bundle):
    code, out, rd = run_bundle("bundle_bnsp_like", "--no-ocr")
    rows = {str(r["Halaman"]): r for r in sheet_rows(rd / "unmatched_report.xlsx", "unmatched")}
    assert rows["3-4"]["Alasan"] == "partial name"
    assert rows["3-4"]["Kandidat 1 NID"] == "0077003"
    assert_accounting(rd)

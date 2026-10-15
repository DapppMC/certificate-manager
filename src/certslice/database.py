from __future__ import annotations

from pathlib import Path
import re

import openpyxl

from certslice.errors import InputError
from certslice.models import Employee
from certslice.normalize import normalize_name


def header_key(v) -> str:
    return re.sub(r"[\s.\-]+", "_", str(v).strip().casefold()) if v is not None else ""


def _find_column_index(keys: list[str], aliases: list[str]) -> int:
    for alias in aliases:
        if alias in keys:
            return keys.index(alias)
    return -1


def load_employees(path: Path, cfg: dict) -> list[Employee]:
    path = Path(path)
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise InputError(f"cannot open database file {path}: {e}")

    try:
        preferred = [p.strip().lower() for p in cfg.get("database", {}).get("preferred_sheets", ["data master"])]
        target_sheet = None
        for name in wb.sheetnames:
            if name.strip().lower() in preferred:
                target_sheet = name
                break
        ws = wb[target_sheet] if target_sheet else wb.worksheets[0]

        max_search = cfg.get("database", {}).get("header_search_rows", 10)
        nid_aliases = cfg.get("database", {}).get("nid", ["nid"])
        nama_aliases = cfg.get("database", {}).get("nama", ["nama"])
        email_aliases = cfg.get("database", {}).get("email", ["email"])
        unit_aliases = cfg.get("database", {}).get("unit", ["unit"])
        jabatan_aliases = cfg.get("database", {}).get("jabatan", ["jabatan"])

        header_row_idx = None
        nid_idx = -1
        nama_idx = -1
        email_idx = -1
        unit_idx = -1
        jabatan_idx = -1

        rows = list(ws.iter_rows(values_only=True))
        scanned_keys_list: list[list[str]] = []

        for r_idx, row in enumerate(rows[:max_search]):
            row_keys = [header_key(c) for c in row]
            scanned_keys_list.append(row_keys)
            c_nid = _find_column_index(row_keys, nid_aliases)
            c_nama = _find_column_index(row_keys, nama_aliases)
            if c_nid != -1 and c_nama != -1:
                header_row_idx = r_idx
                nid_idx = c_nid
                nama_idx = c_nama
                email_idx = _find_column_index(row_keys, email_aliases)
                unit_idx = _find_column_index(row_keys, unit_aliases)
                jabatan_idx = _find_column_index(row_keys, jabatan_aliases)
                break

        if header_row_idx is None:
            nid_seen = any(_find_column_index(k, nid_aliases) != -1 for k in scanned_keys_list)
            nama_seen = any(_find_column_index(k, nama_aliases) != -1 for k in scanned_keys_list)
            if not nid_seen and not nama_seen:
                raise InputError("database is missing the required column(s): nid, nama")
            elif not nid_seen:
                raise InputError("database is missing the required column(s): nid")
            else:
                raise InputError("database is missing the required column(s): nama")

        employees: list[Employee] = []
        for row in rows[header_row_idx + 1:]:
            if len(row) <= max(nid_idx, nama_idx):
                continue
            nid_raw = row[nid_idx]
            nama_raw = row[nama_idx]
            if nid_raw is None or nama_raw is None:
                continue

            if isinstance(nid_raw, int):
                nid = str(nid_raw)
            elif isinstance(nid_raw, float) and nid_raw.is_integer():
                nid = str(int(nid_raw))
            else:
                nid = str(nid_raw).strip()

            nama = str(nama_raw).strip()
            if not nid or not nama:
                continue

            email = ""
            if email_idx != -1 and len(row) > email_idx and row[email_idx] is not None:
                email = str(row[email_idx]).strip()

            unit = ""
            if unit_idx != -1 and len(row) > unit_idx and row[unit_idx] is not None:
                unit = str(row[unit_idx]).strip()

            jabatan = ""
            if jabatan_idx != -1 and len(row) > jabatan_idx and row[jabatan_idx] is not None:
                jabatan = str(row[jabatan_idx]).strip()

            employees.append(
                Employee(
                    nid=nid,
                    nama=nama,
                    email=email,
                    unit=unit,
                    jabatan=jabatan,
                    norm=normalize_name(nama),
                )
            )

        if not employees:
            raise InputError("database has no employee rows below the header")

        return employees
    finally:
        wb.close()

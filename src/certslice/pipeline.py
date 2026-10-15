from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
import json
from pathlib import Path
import re

import openpyxl
import pymupdf

from certslice.config import load_config
from certslice.errors import InputError


@dataclass
class RunResult:
    run_dir: Path | None
    exit_code: int
    summary: dict


def _fold_text(s: str) -> str:
    s = s.lower()
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return " ".join(s.split())


def run(
    pdf_path,
    db_path,
    out_dir,
    *,
    dry_run=False,
    pages_per_cert="auto",
    ocr=True,
    config_path=None,
    reader=None,
) -> RunResult:
    cfg = load_config(config_path)
    pdf_file = Path(pdf_path)
    db_file = Path(db_path)
    out_path = Path(out_dir)

    # (a) Check file existence
    if not pdf_file.exists():
        raise InputError(f"PDF file not found: {pdf_file}")
    if not db_file.exists():
        raise InputError(f"Database file not found: {db_file}")

    # (b) Read database
    try:
        wb = openpyxl.load_workbook(db_file, read_only=True, data_only=True)
    except Exception as e:
        raise InputError(f"Could not open database file: {e}")

    target_sheet = None
    for name in wb.sheetnames:
        if name.strip().lower() == "data master":
            target_sheet = name
            break
    ws = wb[target_sheet] if target_sheet else wb.worksheets[0]

    header_row_idx = None
    nid_col = -1
    name_col = -1
    email_col = -1
    unit_col = -1

    rows = list(ws.iter_rows(values_only=True))
    for r_idx, row in enumerate(rows[:10]):
        cells = [str(v).strip().lower() if v is not None else "" for v in row]
        if "nid" in cells:
            header_row_idx = r_idx
            nid_col = cells.index("nid")
            for c_idx, val in enumerate(cells):
                if val in ("nama_lengkap", "nama lengkap", "nama"):
                    name_col = c_idx
                elif "email" in val:
                    email_col = c_idx
                elif "unit" in val:
                    unit_col = c_idx
            break

    if header_row_idx is None or nid_col == -1 or name_col == -1:
        wb.close()
        raise InputError("Database missing required NID or name columns")

    employees = []
    for row in rows[header_row_idx + 1:]:
        if len(row) <= max(nid_col, name_col):
            continue
        nid_raw = row[nid_col]
        nama_raw = row[name_col]
        if nid_raw is None or nama_raw is None:
            continue

        if isinstance(nid_raw, int):
            nid = str(nid_raw)
        elif isinstance(nid_raw, float) and nid_raw.is_integer():
            nid = str(int(nid_raw))
        else:
            nid = str(nid_raw).strip()

        nama = str(nama_raw).strip()
        if not nid or not nama:
            continue

        email = ""
        if email_col != -1 and len(row) > email_col and row[email_col] is not None:
            email = str(row[email_col]).strip()

        unit = ""
        if unit_col != -1 and len(row) > unit_col and row[unit_col] is not None:
            unit = str(row[unit_col]).strip()

        employees.append({
            "nid": nid,
            "nama": nama,
            "email": email,
            "unit": unit,
        })
    wb.close()

    if not employees:
        raise InputError("Database contains no employee rows")

    # (c) Open PDF and read pages
    try:
        doc = pymupdf.open(pdf_file)
    except Exception as e:
        raise InputError(f"Could not open PDF file: {e}")

    num_pages = doc.page_count
    if num_pages == 0:
        doc.close()
        raise InputError("PDF contains no pages")

    page_records = []
    groups = []

    # (d) & (e) Simple match and grouping
    for page_idx in range(num_pages):
        page_num = page_idx + 1
        page = doc[page_idx]

        if reader is not None:
            page_text_obj = reader.read(page)
            raw_text = page_text_obj.text
            source = page_text_obj.source
        else:
            raw_text = page.get_text("text")
            source = "text"

        folded_page = " " + _fold_text(raw_text) + " "

        matched_employees = []
        for emp in employees:
            folded_name = " " + _fold_text(emp["nama"]) + " "
            if folded_name in folded_page:
                matched_employees.append(emp)

        distinct_nids = {emp["nid"] for emp in matched_employees}

        if len(distinct_nids) == 1:
            matched_emp = matched_employees[0]
            status = "accepted"
            reason = ""
            file_name = f"{matched_emp['nid']}_{matched_emp['nama']}.pdf"
            cand_page = [{
                "nid": matched_emp["nid"],
                "nama": matched_emp["nama"],
                "score": 100.0,
            }]
            cand_group = [{
                "nid": matched_emp["nid"],
                "nama": matched_emp["nama"],
                "unit": matched_emp["unit"],
                "score": 100.0,
            }]
            group_entry = {
                "id": page_num,
                "pages": [page_num],
                "status": status,
                "reason": reason,
                "nid": matched_emp["nid"],
                "nama": matched_emp["nama"],
                "email": matched_emp["email"],
                "unit": matched_emp["unit"],
                "candidates": cand_group,
                "fields": {},
                "expiry_source": None,
                "renewal": None,
                "file_name": file_name,
            }
        else:
            status = "unmatched"
            reason = "no confident holder"
            cand_page = []
            group_entry = {
                "id": page_num,
                "pages": [page_num],
                "status": status,
                "reason": reason,
                "nid": None,
                "nama": None,
                "email": None,
                "unit": None,
                "candidates": [],
                "fields": {},
                "expiry_source": None,
                "renewal": None,
                "file_name": None,
            }

        page_records.append({
            "page": page_num,
            "source": source,
            "candidates": cand_page,
        })
        groups.append(group_entry)

    # (f) Create run folder
    out_path.mkdir(parents=True, exist_ok=True)
    base_name = f"run-{datetime.now().strftime('%Y%m%d-%H%M%S')}"
    run_dir = out_path / base_name
    counter = 2
    while run_dir.exists():
        run_dir = out_path / f"{base_name}-{counter}"
        counter += 1
    run_dir.mkdir(parents=True, exist_ok=True)

    # (g) Write PDFs unless dry_run
    if not dry_run:
        pdf_dir = run_dir / "pdf"
        pdf_dir.mkdir(parents=True, exist_ok=True)
        for grp in groups:
            if grp["status"] == "accepted":
                p_idx = grp["pages"][0] - 1
                new_doc = pymupdf.open()
                new_doc.insert_pdf(doc, from_page=p_idx, to_page=p_idx)
                new_doc.save(pdf_dir / grp["file_name"])
                new_doc.close()

    doc.close()

    # (h) Write plan.json
    plan_data = {
        "version": 1,
        "input_pdf": str(pdf_file),
        "database": str(db_file),
        "page_count": num_pages,
        "dry_run": dry_run,
        "pages_per_cert": pages_per_cert,
        "template_names": [],
        "pages": page_records,
        "groups": groups,
        "timings": {
            "read_seconds": 0.0,
            "total_seconds": 0.0,
            "ocr_pages": 0,
            "ocr_seconds_per_page": None,
        },
    }

    plan_path = run_dir / "plan.json"
    plan_path.write_text(json.dumps(plan_data, ensure_ascii=False, indent=2), encoding="utf-8")

    # (i) Summary and return
    accepted_count = sum(1 for g in groups if g["status"] == "accepted")
    unmatched_count = sum(1 for g in groups if g["status"] == "unmatched")
    summary = {
        "pages": num_pages,
        "accepted": accepted_count,
        "unmatched": unmatched_count,
    }

    return RunResult(run_dir=run_dir, exit_code=0, summary=summary)

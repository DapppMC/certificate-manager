"""Unified processing and bulk renaming module for Certificate Slicer.

Handles single-certificate PDFs, multi-page bundle slicing, transmittal letter (roster)
detection and parsing, and unified 16-column Excel report generation.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
import concurrent.futures
import gc
import json
import logging
import os
from pathlib import Path
import re
import shutil
import threading
from typing import Any, Callable

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pymupdf
from rapidfuzz import fuzz

from certslice.config import load_config
from certslice.database import load_employees
from certslice.fields import extract_dates, extract_title
from certslice.gemini import GeminiRescue, load_user_settings
from certslice.grouping import PageInfo, group_pages
from certslice.matcher import cue_filter, decide, page_candidates
from certslice.models import Candidate, Employee
from certslice.naming import clean_field
from certslice.normalize import fold, name_tokens
from certslice.ocr import OcrEngine
from certslice.pipeline import extract_raw_holder_name

logger = logging.getLogger(__name__)


@dataclass
class CertJob:
    job_id: int
    src_pdf: Path
    pages: list[int]  # 0-indexed page numbers
    is_bundle: bool = False


@dataclass
class BulkItem:
    index: int
    src_path: Path
    file_name: str = ""
    nid: str = ""
    nama: str = ""
    jabatan: str = ""
    unit: str = ""
    is_skp: bool = False
    level: str = ""
    peran: str = ""
    judul: str = ""
    number: str = ""
    issue_date: str = ""
    expiry_date: str = ""
    status: str = "accepted"  # accepted | unmatched
    reason: str = ""
    candidates: list[dict[str, Any]] = field(default_factory=list)
    pages: list[int] = field(default_factory=lambda: [0])  # 0-indexed pages in src_path
    is_bundle: bool = False


@dataclass
class BulkResult:
    run_dir: Path
    total_files: int
    accepted_count: int
    review_count: int
    items: list[BulkItem]


def is_roster_doc(pdf_path: Path, ocr_engine: OcrEngine | None = None) -> bool:
    """Detect if a PDF is a transmittal letter / roster reference table rather than an employee certificate."""
    name_lower = pdf_path.stem.lower()
    roster_name_keywords = [
        "surat pengantar", "lampiran", "roster", "daftar penerima", "pengantar pengiriman",
        "daftar peserta", "daftar sertifikat", "rekap sertifikat", "pengiriman sertifikasi",
        "keterangan pengiriman", "pengiriman sertifikat", "hasil perpanjangan", "perpanjangan kemnaker"
    ]
    if re.search(r"[-_ ]1-4$", name_lower) or any(k in name_lower for k in roster_name_keywords):
        return True
    try:
        doc = pymupdf.open(pdf_path)
        has_any_text = False
        roster_text_keywords = [
            "LAMPIRAN SURAT PENGANTAR",
            "SURAT PENGANTAR PENGIRIMAN",
            "SURAT KETERANGAN PENGIRIMAN",
            "KODE OKUPASI JABATAN",
            "PENERBITANSERTIFIKASI",
            "PENERBITAN SERTIFIKASI",
            "DAFTAR PENERIMA",
            "DAFTAR SERTIFIKASI",
            "HASIL PERPANJANGAN",
            "HASIL PERPANJANGAN KEMNAKER",
        ]
        for p in range(min(3, len(doc))):
            t = doc[p].get_text()
            if t.strip():
                has_any_text = True
                t_upper = t.upper()
                if any(k in t_upper for k in roster_text_keywords) or (
                    "KODE OKUPASI" in t_upper and ("NO SERTIFIKAT" in t_upper or "SERTIFIKAT" in t_upper)
                ) or ("JENIS SERTIFIKASI" in t_upper and ("LISENSI" in t_upper or "SKP" in t_upper)):
                    doc.close()
                    return True

        # If document has NO embedded digital text (scanned bitmap), ONLY perform OCR if filename has roster cues
        if not has_any_text and ocr_engine:
            roster_filename_cues = ["surat", "lampiran", "roster", "daftar", "pengantar", "pengiriman", "rekap", "tabel", "skp", "kemnaker", "1-4", "penerima"]
            if any(k in name_lower for k in roster_filename_cues):
                for p in range(min(2, len(doc))):
                    t = ocr_engine.page_text(doc[p])
                    t_upper = t.upper()
                    if any(k in t_upper for k in roster_text_keywords) or (
                        "KODE OKUPASI" in t_upper and ("NO SERTIFIKAT" in t_upper or "SERTIFIKAT" in t_upper)
                    ) or ("JENIS SERTIFIKASI" in t_upper and ("LISENSI" in t_upper or "SKP" in t_upper)):
                        doc.close()
                        return True
        doc.close()
    except Exception:
        pass
    return False


def parse_roster_doc(pdf_path: Path, ocr_engine: OcrEngine) -> list[dict[str, Any]]:
    """Parse transmittal letter / roster PDF table to extract employee names, levels, role names, and cert numbers."""
    doc = pymupdf.open(pdf_path)
    entries: list[dict[str, Any]] = []

    for p in range(len(doc)):
        t = doc[p].get_text()
        if not t.strip() or len(t.splitlines()) < 5:
            t = ocr_engine.page_text(doc[p])
        lines = [l.strip() for l in t.splitlines() if l.strip()]

        # Check if table page (must have KUALIFIKASI or certificate number patterns)
        kual_indices = [i for i, l in enumerate(lines) if "KUALIFIKASI" in l]
        if not kual_indices:
            kual_indices = [
                i for i, l in enumerate(lines)
                if re.search(r"(\d{4}\.[\d\.]+\.[A-Z0-9]+\.\d{2}\.\d{4})", l)
            ]
        if not kual_indices:
            # Check if Kemnaker / simple numbered roster table
            is_kemnaker = any("JENIS SERTIFIKASI" in l.upper() for l in lines) or any("PERPANJANGAN KEMNAKER" in l.upper() for l in lines) or any("PENGIRIMAN SERTIFIKAT" in l.upper() for l in lines)
            if is_kemnaker:
                start_table = False
                for idx, l in enumerate(lines):
                    if any(hdr in l.upper() for hdr in ["JENIS SERTIFIKASI", "NAMA", "PERUSAHAAN"]):
                        start_table = True
                        continue
                    if not start_table:
                        continue

                    m = re.match(r"^(\d{1,2})\s+([A-Za-z\.\'\s]+)$", l)
                    if m and int(m.group(1)) > 0:
                        num = int(m.group(1))
                        name = m.group(2).strip()
                        entries.append({
                            "row": num,
                            "nama": name,
                            "level": "",
                            "cert_no": "",
                            "reg_no": "",
                            "kode_okupasi": "",
                            "peran": "",
                            "tokens": [name],
                            "lines": [l],
                        })
                    elif any(rk in l for rk in ["Perpanjangan", "Operator", "AK3", "Mutasi"]):
                        role = l
                        if idx + 1 < len(lines) and not re.match(r"^\d{1,2}\s+", lines[idx + 1]) and not any(k in lines[idx + 1] for k in ["Soft File", "PT PLN", "Power", "Services"]):
                            role += " " + lines[idx + 1]
                        if entries:
                            entries[-1]["peran"] = role.strip()
                    elif any(sk in l for sk in ["Soft File", "PT PLN", "Power Services"]):
                        continue

                # Forward-fill role for consecutive rows that share the same certification
                last_r = ""
                for e in entries:
                    if e["peran"]:
                        last_r = e["peran"]
                    elif last_r:
                        e["peran"] = last_r
            continue

        # Header ends before the first table entry
        start_table = 0
        for i, l in enumerate(lines[: kual_indices[0]]):
            if any(k in l.lower() for k in ["pengirim penerima", "kesesuaian", "no sertifikat", "registrasi", "kode okupasi"]):
                start_table = i + 1

        # Footer begins after the last table entry
        end_table = len(lines)
        for i, l in enumerate(lines[kual_indices[-1] :], start=kual_indices[-1]):
            if any(term in l.lower() for term in ["pt skp tenaga listrik", "tebet raya", "021-22831822", "akreditasi kementerian"]):
                end_table = i
                break

        # Calculate row ranges between anchors
        row_ranges: list[tuple[int, int]] = []
        for idx, k_pos in enumerate(kual_indices):
            # Start of row
            if idx == 0:
                r_start = start_table
            else:
                prev_k = kual_indices[idx - 1]
                split_at = None
                for j in range(prev_k + 1, k_pos):
                    m = re.match(r"^(\d{1,2})(?:\s+[A-Za-z].*)?$", lines[j])
                    if m and int(m.group(1)) > 0:
                        split_at = j
                        break
                if split_at is not None:
                    r_start = split_at
                else:
                    cert_pos = None
                    for j in range(prev_k, k_pos):
                        if re.search(r"\d{4}\.[\d\.]+\.[A-Z0-9]+\.\d{2}\.\d{4}", lines[j]):
                            cert_pos = j
                            break
                    if cert_pos is not None and cert_pos + 1 < k_pos:
                        r_start = cert_pos + 1
                    else:
                        r_start = (prev_k + k_pos) // 2 + 1

            # End of row
            if idx == len(kual_indices) - 1:
                r_end = end_table
            else:
                next_k = kual_indices[idx + 1]
                split_at = None
                for j in range(k_pos + 1, next_k):
                    m = re.match(r"^(\d{1,2})(?:\s+[A-Za-z].*)?$", lines[j])
                    if m and int(m.group(1)) > 0:
                        split_at = j
                        break
                if split_at is not None:
                    r_end = split_at
                else:
                    cert_pos = None
                    for j in range(k_pos, next_k):
                        if re.search(r"\d{4}\.[\d\.]+\.[A-Z0-9]+\.\d{2}\.\d{4}", lines[j]):
                            cert_pos = j
                            break
                    if cert_pos is not None and cert_pos + 1 < next_k:
                        r_end = cert_pos + 1
                    else:
                        r_end = (k_pos + next_k) // 2 + 1

            row_ranges.append((r_start, r_end))

        # Extract structured data from each row segment
        for r_start, r_end in row_ranges:
            r_lines = lines[r_start:r_end]
            full_r = " ".join(r_lines)

            m_c = re.search(r"(\d{4}\.[\d\.]+\.[A-Z0-9]+\.\d{2}\.\d{4})", full_r)
            if not m_c:
                m_c = re.search(r"(\d{4}\.[\d\.]+\.[A-Z0-9]+\.[\d\.]+)", full_r)
            cert_no = m_c.group(1) if m_c else ""

            m_k = re.search(r"KUALIFIKASI\s*\.?\s*([1-9])", full_r, re.IGNORECASE)
            m_l = re.search(r"Level\s*([1-9])", full_r, re.IGNORECASE)
            level = f"Level {m_l.group(1)}" if m_l else (f"Level {m_k.group(1)}" if m_k else "")

            m_reg = re.search(r"(\d{5,}\.[\d\.]+)", full_r)
            reg_no = m_reg.group(1) if m_reg else ""

            m_kode = re.search(r"([A-Z]\.\d+\.\d+\.KUALIFIKASI\.[1-9]\.[A-Z0-9]+)", full_r)
            kode_okupasi = m_kode.group(1) if m_kode else ""

            row_id: int | None = None
            clean_tokens: list[str] = []
            role_tokens: list[str] = []
            bad_role_terms = [
                "kualifikasi", "akreditasi", "kementerian", "esdm", "www.ptskp",
                "tebet raya", "skp", "checklist", "kesesuaian", "pengirim", "penerima",
                "gmail.com", "mohon untuk", "021-22831822"
            ]

            for item in r_lines:
                m_lead = re.match(r"^(\d{1,2})\s+(.+)$", item)
                if m_lead and row_id is None:
                    row_id = int(m_lead.group(1))
                    clean_tokens.append(m_lead.group(2))
                elif re.match(r"^(\d{1,2})$", item) and row_id is None:
                    row_id = int(item)
                else:
                    clean_tokens.append(item)

                m_kode_line = re.search(r"[A-Z]\.\d+\.\d+\.KUALIFIKASI\.[1-9]\.[A-Z0-9]+(?:\s+(.*))?", item)
                if m_kode_line and m_kode_line.group(1):
                    role_tokens.append(m_kode_line.group(1))
                elif (
                    not any(b in item.lower() for b in bad_role_terms)
                    and not re.search(r"\d{4}\.[\d\.]+\.[A-Z0-9]+\.\d{2}\.\d{4}", item)
                    and not re.search(r"^\d+$", item)
                    and not re.match(r"^Level\s*[1-9]", item, re.IGNORECASE)
                ):
                    role_keywords = [
                        "manajer", "operator", "supervisor", "pelaksana", "senior",
                        "pemeliharaan", "operasi", "peralatan", "boiler", "turbin",
                        "generator", "plts", "pltu", "pltd", "control room", "assistant",
                        "engineer", "planner", "coal", "handling", "mesin"
                    ]
                    if any(rk in item.lower() for rk in role_keywords):
                        role_tokens.append(item)

            peran = " ".join(role_tokens).strip()

            if cert_no or level:
                entries.append({
                    "row": row_id or len(entries) + 1,
                    "level": level,
                    "cert_no": cert_no,
                    "reg_no": reg_no,
                    "kode_okupasi": kode_okupasi,
                    "peran": peran,
                    "tokens": clean_tokens,
                    "lines": r_lines,
                })

    doc.close()
    return entries


def _enrich_roster_entry_role(entry: dict[str, Any], emp_nama: str) -> dict[str, Any]:
    """Extract peran and matched name tokens from entry using employee name words."""
    folded_target = fold(emp_nama)
    target_words = set(folded_target.split())
    matched_words: list[str] = []
    role_words: list[str] = []
    bad_role_terms = [
        "kualifikasi", "akreditasi", "kementerian", "esdm", "www.ptskp",
        "tebet raya", "skp", "advancing", "checklist", "kesesuaian",
        "pengirim", "penerima", "pard", "gmail.com", "mohon untuk",
        "jabatan", "nama", "registrasi", "kode okupasi"
    ]
    for tok in entry.get("tokens", []):
        ft = fold(tok)
        if any(w in ft.split() for w in target_words):
            matched_words.append(tok)
        elif not any(b in ft for b in bad_role_terms) and not re.search(r"[A-Z]\.\d+\.\d+\.KUALIFIKASI", tok):
            role_words.append(tok)
    res = dict(entry)
    if role_words and not res.get("peran"):
        res["peran"] = " ".join(role_words).strip()
    if matched_words:
        res["nama"] = " ".join(matched_words).strip()
    return res


def match_roster(emp_nama: str, entries: list[dict[str, Any]], cert_no: str | None = None) -> dict[str, Any] | None:
    """Find matching employee in roster entries by Certificate Number (exact) or Employee Name (fuzzy)."""
    if not entries:
        return None

    # 1. Primary: Exact match by Certificate Number
    if cert_no:
        clean_target_cert = cert_no.strip()
        for e in entries:
            if e.get("cert_no") and e["cert_no"].strip() == clean_target_cert:
                res = dict(e)
                if not res.get("peran") and emp_nama:
                    res = _enrich_roster_entry_role(res, emp_nama)
                return res

    # 2. Fallback: Fuzzy match by Employee Name tokens
    if not emp_nama:
        return None

    folded_target = fold(emp_nama)
    target_words = set(folded_target.split())
    best_entry: dict[str, Any] | None = None
    best_score = 0.0

    for e in entries:
        token_str = " ".join(e["tokens"])
        folded_tokens = fold(token_str)
        row_words = set(folded_tokens.split())
        inter = target_words.intersection(row_words)
        if len(inter) == len(target_words) and len(target_words) > 0:
            score = 100.0 + len(inter)
        elif len(inter) > 0:
            score = (len(inter) / len(target_words)) * 90.0
        else:
            score = 0.0
        if score > best_score and score >= 60.0:
            best_score = score
            best_entry = e

    if best_entry:
        return _enrich_roster_entry_role(dict(best_entry), emp_nama)
    return None


UNIT_STOP_WORDS = {
    'pltu', 'pltd', 'plts', 'pltg', 'pltmg', 'plta', 'pt', 'pln', 'nps', 'kp', 'up', 'np',
    'sto', 'cng', 'biomassa', 'jawa', 'dan', 'tenaga', 'listrik', 'unit', 'updk', 'sektor',
    'kantor', 'pusat'
}

JABATAN_STOP_WORDS = {
    'pt', 'pln', 'nps', 'dan', 'di', 'pada', 'staf', 'staff', 'pegawai', 'karyawan', 'bagian', 'bidang', 'divisi', 'kantor'
}

EXTENDED_CONTINUATION_KEYWORDS = [
    "rincian unit", "details for unit", "transkrip", "materi",
    "jam pelajaran", "daftar unit", "kompetensi inti", "unit kompetensi",
    "silabus", "kurikulum", "fokus pengembangan", "jadwal", "mata pelajaran",
    "struktur program", "daftar modul", "hasil evaluasi"
]


def get_unit_keywords(unit_str: str) -> set[str]:
    """Extract distinctive location keywords from a company unit name."""
    if not unit_str:
        return set()
    toks = fold(unit_str).split()
    return {t for t in toks if t not in UNIT_STOP_WORDS and len(t) >= 4}


def get_jabatan_keywords(jab_str: str) -> set[str]:
    """Extract distinctive role keywords from an employee position/jabatan."""
    if not jab_str:
        return set()
    toks = fold(jab_str).split()
    return {t for t in toks if t not in JABATAN_STOP_WORDS and len(t) >= 4}


def normalize_digraphs(s: str) -> str:
    """Normalize common Indonesian spelling digraphs and repeated consonants."""
    s = s.replace("kh", "h").replace("ch", "h").replace("sy", "s").replace("dh", "d")
    s = s.replace("oe", "u").replace("ie", "i").replace("th", "t").replace("ph", "f")
    return re.sub(r"(.)\1+", r"\1", s)


def token_similarity(qt: str, et: str) -> float:
    """Robust token similarity for Indonesian names with typo tolerance and prefix matching."""
    if qt == et:
        return 100.0
    if len(qt) >= 3 and len(et) >= 3 and (et.startswith(qt) or qt.startswith(et)):
        return 95.0
    r = float(fuzz.ratio(qt, et))
    if r >= 80.0:
        return r
    if len(qt) >= 4 and len(et) >= len(qt):
        r_pref = float(fuzz.ratio(qt, et[:len(qt)]))
        if r_pref >= 80.0:
            return r_pref
    if len(qt) >= 4 and len(et) >= 4:
        r_norm = float(fuzz.ratio(normalize_digraphs(qt), normalize_digraphs(et)))
        if r_norm >= 80.0:
            return r_norm
        if len(et) >= len(qt):
            r_norm_pref = float(fuzz.ratio(normalize_digraphs(qt), normalize_digraphs(et[:len(qt)])))
            if r_norm_pref >= 80.0:
                return r_norm_pref
    return 0.0


def expand_name_abbreviations(name: str) -> list[str]:
    """Expand common Indonesian name abbreviations (M., Moh., Muh., Ach., A.)."""
    res = [name]
    folded = fold(name).strip()
    m_match = re.match(r"^(?:m\.?|moh\.?|muh\.?)\s+(.+)$", folded)
    if m_match:
        rest = m_match.group(1)
        res.append(f"muhammad {rest}")
        res.append(f"mochammad {rest}")
        res.append(rest)
    ach_match = re.match(r"^(?:ach\.?|ah\.?|a\.)\s+(.+)$", folded)
    if ach_match:
        rest = ach_match.group(1)
        res.append(f"achmad {rest}")
        res.append(f"ahmad {rest}")
        res.append(f"abdul {rest}")
        res.append(rest)
    return list(dict.fromkeys(res))


def is_auxiliary_doc(pdf_path: Path, ocr_engine: OcrEngine | None = None) -> tuple[bool, str]:
    """Detect if a PDF is an auxiliary non-certificate document (attendance list, invoice, task letter, rundown)."""
    name_lower = pdf_path.stem.lower().replace("_", " ").replace("-", " ")
    aux_name_map = [
        (["absen", "presensi", "daftar hadir", "attendance"], "Daftar Hadir"),
        (["invoice", "kwitansi", "receipt", "bukti bayar", "tagihan"], "Kwitansi / Invoice"),
        (["rundown", "jadwal", "agenda kegiatan", "susunan acara"], "Jadwal / Rundown"),
        (["surat tugas", "spt", "surat permohonan", "berita acara", "bast"], "Surat Tugas / Berita Acara"),
        (["evaluasi", "feedback", "kuesioner"], "Form Evaluasi"),
    ]
    for keywords, cat in aux_name_map:
        if any(k in name_lower for k in keywords):
            return True, cat

    try:
        doc = pymupdf.open(pdf_path)
        if len(doc) > 0:
            t1 = doc[0].get_text()[:1000].lower()
            doc.close()
            for keywords, cat in aux_name_map:
                if any(k in t1 for k in keywords):
                    if "sertifikat" not in t1 and "certificate" not in t1 and "kompeten" not in t1:
                        return True, cat
        else:
            doc.close()
    except Exception:
        pass
    return False, ""


def find_bulk_candidates(
    lines: list[str],
    employees: list[Employee],
    cfg: dict[str, Any],
    stem_cue: str = "",
) -> tuple[str, Candidate | None, str, list[Candidate]]:
    """Resolve employee candidate for a certificate using multi-source cues (text lines, stem, and fuzzy/subset)."""
    doc_text = fold(" ".join(lines) + " " + stem_cue)

    raw_cands = page_candidates(lines, employees, cfg)
    clean_cands: list[Candidate] = []
    for c in raw_cands:
        line_txt = lines[c.line].lower() if 0 <= c.line < len(lines) else ""
        is_date_line = any(dk in line_txt for dk in ["pada tanggal", "tanggal", "date of", "ditetapkan di", "lahir", "birth", "kambuno", "jakarta"])
        if is_date_line and c.score < 95.0:
            continue
        clean_cands.append(c)

    status, top, reason = decide(clean_cands, employees, cfg)
    if status == "accepted" and top is not None:
        if stem_cue:
            clean_stem = stem_cue.replace('_', ' ').replace('-', ' ').replace('.', ' ').strip().lower()
            emp_name = employees[top.index].norm
            stem_toks = [t for t in name_tokens(clean_stem) if len(t) >= 3]
            if stem_toks and not any(st in emp_name for st in stem_toks):
                combined_matches = [
                    (idx, emp) for idx, emp in enumerate(employees)
                    if all(st in emp.norm for st in stem_toks) and any(et in emp.norm for et in name_tokens(lines[top.line] if 0 <= top.line < len(lines) else "") if len(et) >= 3)
                ]
                if len(combined_matches) == 1:
                    c_idx = combined_matches[0][0]
                    c_top = Candidate(index=c_idx, score=95.0, line=top.line, partial=False)
                    return "accepted", c_top, "", [c_top] + clean_cands
        return status, top, reason, clean_cands

    queries: list[str] = []
    if stem_cue:
        clean_stem = stem_cue.replace('_', ' ').replace('-', ' ').replace('.', ' ')
        clean_stem = ''.join(c for c in clean_stem if c.isalnum() or c.isspace()).strip()
        if clean_stem and not any(k in clean_stem.lower() for k in ["sertifikat", "lampiran", "surat", "roster", "certificate", "skp"]):
            queries.append(clean_stem)

    for i, l in enumerate(lines):
        fl = fold(l)
        # Check inverted comma name format, e.g. "Pratama, Wenedi Surya"
        if "," in l and len(l.split(",")) == 2:
            p_last, p_first = [part.strip() for part in l.split(",", 1)]
            if len(p_last.split()) <= 2 and len(p_first.split()) <= 3:
                queries.append(f"{p_first} {p_last}".strip())

        if any(h in fl for h in ["presented to", "diberikan kepada", "sertifikat ini", "menyatakan bahwa"]):
            if i + 1 < len(lines):
                queries.append(lines[i + 1].strip())
        elif len(fl.split()) in (1, 2, 3) and not any(
            k in fl for k in [
                "certificate", "completion", "english", "training", "januari", "februari",
                "maret", "april", "mei", "juni", "juli", "agustus", "september",
                "oktober", "november", "desember", "pt", "kompetensi", "direktur", "executive",
                "presented to", "for successfully", "is able to apply", "the knowledge", "lokasi"
            ]
        ):
            queries.append(l.strip())

    # Strip patronymic words (bin, binti, ibnu) to test base name
    for q in list(queries):
        cleaned_pat = re.sub(r"\b(?:bin|binti|ibnu)\b.*$", "", q, flags=re.IGNORECASE).strip()
        if cleaned_pat and cleaned_pat != q and len(cleaned_pat.split()) >= 1:
            queries.append(cleaned_pat)

    expanded_queries = []
    for q in queries:
        expanded_queries.extend(expand_name_abbreviations(q))
    queries = list(dict.fromkeys(expanded_queries))

    extra_cands: list[Candidate] = []
    seen_indices = set()

    # If clean_cands has only 1 partial candidate (e.g. Two Okta), upgrade it if unique
    if len(clean_cands) == 1 and clean_cands[0].score >= 90.0:
        c = clean_cands[0]
        extra_cands.append(Candidate(index=c.index, score=93.0, line=c.line, partial=False))
        seen_indices.add(c.index)
    else:
        seen_indices = {c.index for c in clean_cands}

    # Combined tokens across queries (e.g. stem 'daniel' + ocr 'ferry')
    combined_query_toks = set()
    for q in queries:
        for t in name_tokens(q):
            if len(t) >= 3:
                combined_query_toks.add(t)

    if len(combined_query_toks) >= 2:
        multi_matches: list[tuple[int, int]] = []
        for idx, emp in enumerate(employees):
            e_toks = set(emp.norm.split())
            common = combined_query_toks.intersection(e_toks)
            if len(common) >= 2:
                multi_matches.append((idx, len(common)))
        if len(multi_matches) == 1:
            idx = multi_matches[0][0]
            if idx not in seen_indices:
                extra_cands.append(Candidate(index=idx, score=94.0, line=0, partial=False))
                seen_indices.add(idx)

    for q in queries:
        q_toks = [t for t in name_tokens(q) if len(t) >= 2]
        if not q_toks:
            continue

        emp_scores: list[tuple[int, float, bool]] = []
        for idx, emp in enumerate(employees):
            e_toks = emp.norm.split()
            matched_sims = []
            is_all_exact = True
            for qt in q_toks:
                best_sim = 0.0
                for et in e_toks:
                    sim = token_similarity(qt, et)
                    if sim > best_sim:
                        best_sim = sim
                if best_sim >= 80.0:
                    matched_sims.append(best_sim)
                    if best_sim < 100.0:
                        is_all_exact = False
                else:
                    break

            if len(matched_sims) == len(q_toks):
                avg_sim = sum(matched_sims) / len(matched_sims)
                emp_scores.append((idx, avg_sim, is_all_exact))

        if emp_scores:
            emp_scores.sort(key=lambda x: -x[1])
            best_sim = emp_scores[0][1]
            top_matches = [x for x in emp_scores if x[1] >= best_sim - 2.0]

            if len(top_matches) == 1:
                idx, sim, is_exact = top_matches[0]
                if idx not in seen_indices:
                    sc = 93.0 if sim >= 88.0 else 90.0
                    extra_cands.append(Candidate(index=idx, score=sc, line=0, partial=not (sim >= 88.0)))
                    seen_indices.add(idx)
            else:
                for idx, sim, is_exact in top_matches:
                    if idx not in seen_indices:
                        extra_cands.append(Candidate(index=idx, score=min(90.0, sim), line=0, partial=True))
                        seen_indices.add(idx)

        # Whole string fuzzy match
        q_fold = fold(q)
        for idx, emp in enumerate(employees):
            if idx in seen_indices:
                continue
            r = fuzz.token_sort_ratio(q_fold, emp.norm)
            if r >= 85:
                extra_cands.append(Candidate(index=idx, score=float(r), line=0, partial=True))
                seen_indices.add(idx)

    all_cands = clean_cands + extra_cands

    # Contextual Unit & Jabatan Boosting
    unit_boosted_cands = []
    for c in all_cands:
        emp = employees[c.index]
        u_kw = get_unit_keywords(emp.unit)
        j_kw = get_jabatan_keywords(emp.jabatan)
        has_unit_match = bool(u_kw and any(kw in doc_text for kw in u_kw))
        has_jab_match = bool(j_kw and any(kw in doc_text for kw in j_kw))

        boost = 0.0
        if has_unit_match:
            boost += 15.0
        if has_jab_match:
            boost += 8.0

        if boost > 0.0:
            boosted_score = c.score + boost
            unit_boosted_cands.append(Candidate(index=c.index, score=boosted_score, line=c.line, partial=False))
        else:
            unit_boosted_cands.append(c)

    unit_boosted_cands.sort(key=lambda c: (-c.score, c.index))

    if unit_boosted_cands:
        best = unit_boosted_cands[0]
        is_clear_winner = False
        if len(unit_boosted_cands) == 1:
            is_clear_winner = (best.score >= 92.0 and not best.partial)
        else:
            second = unit_boosted_cands[1]
            is_clear_winner = (best.score >= 92.0 and best.score - second.score >= 3.0)

        if is_clear_winner:
            return "accepted", best, "", unit_boosted_cands
        else:
            return "unmatched", best, "perlu konfirmasi (beberapa kandidat mirip di database)", unit_boosted_cands

    return "unmatched", None, f"Nama '{queries[0] if queries else stem_cue}' tidak ditemukan di database master", []


def _detect_skp_offline(full_text: str) -> tuple[bool, str, str]:
    """Detect if text is from an SKP certificate and extract Level & Peran offline."""
    upper = full_text.upper()
    is_skp = any(kw in upper for kw in [
        "SKP TENAGA LISTRIK", "PT. SKP", "PT.SERTIFIKASIKOMPETENSI",
        "PEMBANGKITAN TENAGA LISTRIK", "DJL.4", "KOMPETENSI INTI"
    ])
    level = ""
    peran = ""
    if not is_skp:
        return False, "", ""

    # Try extracting level
    lvl_match = re.search(
        r"bagi\s+(Pelaksana\s+[A-Za-z]+|Operator\s+[A-Za-z]+|Teknisi\s+[A-Za-z]+|Assistant\s+Engineer|Junior\s+Analyst|Ahli\s+[A-Za-z]+)",
        full_text,
        re.IGNORECASE,
    )
    if lvl_match:
        level = lvl_match.group(1).strip()
    else:
        lvl_num = re.search(r"Level\s*([1-9])", full_text, re.IGNORECASE)
        if lvl_num:
            level = f"Level {lvl_num.group(1)}"
        else:
            lvl_jenj = re.search(r"KUALIFIKASI\s*\.?\s*([1-9])", full_text, re.IGNORECASE)
            if lvl_jenj:
                level = f"Level {lvl_jenj.group(1)}"

    # Try extracting peran
    peran_match = re.search(
        r"(?:^\s*|\n)\s*(?:\d+\.\s*)?Unit Kompetensi[^\n]*\n\s*[:]?\s*(.+?)(?:\s+bagi|\s*\(Operating|\n|$)",
        full_text,
        re.IGNORECASE,
    )
    if not peran_match:
        peran_match = re.search(r"Unit Kompetensi\s*[:]\s*(.+?)(?:\s*bagi|\(Operating|\n|$)", full_text, re.IGNORECASE)

    if peran_match:
        cand_p = peran_match.group(1).strip(" :-_.")
        if cand_p and not any(bad in cand_p.lower() for bad in ["kompetensi inti", "rincian unit", "unit competency", "details for", "/profesi", "occupational"]):
            peran = cand_p

    if not peran:
        jab_match = re.search(
            r"Jabatan/Profesi[^\n]*\n\s*[:]?\s*(.+?)(?:\s*\(|\n|$)",
            full_text,
            re.IGNORECASE,
        )
        if not jab_match:
            jab_match = re.search(r"Jabatan\s*[:]\s*(.+?)(?:\n|$)", full_text, re.IGNORECASE)
        if jab_match:
            cand_j = jab_match.group(1).strip(" :-_.")
            if cand_j and not any(bad in cand_j.lower() for bad in ["profesi", "occupational", "description", "deskripsi", "/"]):
                peran = cand_j

    return True, level, peran


_cache_lock = threading.Lock()


def get_optimal_worker_count() -> int:
    """Calculate safe, balanced worker count based on CPU cores."""
    cpus = os.cpu_count() or 2
    if cpus <= 4:
        return 2  # Safe for dual/quad core low-spec laptops
    elif cpus <= 8:
        return 3
    else:
        return 4  # Sweet spot for 8+ cores (e.g. Ryzen 7) without RAM thrashing


def _get_page_text_cached(
    doc: pymupdf.Document,
    p_idx: int,
    pdf_path: Path,
    ocr_engine: OcrEngine | None,
    text_cache: dict[tuple[str, int], str],
) -> str:
    key = (str(pdf_path.resolve()), p_idx)
    with _cache_lock:
        if key in text_cache:
            return text_cache[key]

    p = doc[p_idx]
    t = p.get_text()
    if not t.strip() and ocr_engine:
        t = ocr_engine.page_text(p)

    with _cache_lock:
        text_cache[key] = t
    return t


def prepare_cert_jobs(
    pdf_paths: list[Path],
    employees: list[Employee],
    cfg: dict[str, Any],
    ocr_engine: OcrEngine,
    pages_per_cert: str = "auto",
    text_cache: dict[tuple[str, int], str] | None = None,
) -> tuple[list[CertJob], list[Path], list[dict[str, Any]], list[tuple[Path, str]]]:
    """Scan input PDFs: separate roster docs and auxiliary attachments, detect 1/2/3-page certs vs bundles, and build CertJobs."""
    if text_cache is None:
        text_cache = {}
    roster_paths: list[Path] = []
    roster_entries: list[dict[str, Any]] = []
    aux_paths: list[tuple[Path, str]] = []
    cert_pdfs: list[Path] = []

    # 1. Separate Roster Docs & Auxiliary Documents (Attendance, Invoices, Rundowns, etc.)
    for p in pdf_paths:
        if is_roster_doc(p, ocr_engine):
            roster_paths.append(p)
            try:
                entries = parse_roster_doc(p, ocr_engine)
                roster_entries.extend(entries)
                logger.info("Dokumen Roster terdeteksi: %s (%d penerima)", p.name, len(entries))
            except Exception as e:
                logger.warning("Gagal mem-parse roster %s: %s", p.name, e)
        else:
            is_aux, aux_cat = is_auxiliary_doc(p, ocr_engine)
            if is_aux:
                aux_paths.append((p, aux_cat))
                logger.info("Dokumen Lampiran Non-Sertifikat terdeteksi: %s (%s)", p.name, aux_cat)
            else:
                cert_pdfs.append(p)

    jobs: list[CertJob] = []

    # 2. Deconstruct Certificate PDFs
    continuation_patterns = cfg.get("grouping", {}).get("continuation_patterns", [])
    accept_score = cfg.get("matching", {}).get("accept_score", 92)

    for pdf_idx, pdf_path in enumerate(cert_pdfs, 1):
        doc = pymupdf.open(pdf_path)
        n_pages = len(doc)
        if n_pages == 1:
            jobs.append(CertJob(job_id=len(jobs) + 1, src_pdf=pdf_path, pages=[0], is_bundle=False))
        elif n_pages == 2:
            t2_raw = doc[1].get_text()
            if t2_raw.strip():
                t2 = _get_page_text_cached(doc, 1, pdf_path, ocr_engine, text_cache)
                t2_folded = fold(t2).lower()
                is_cont = (
                    any(re.search(pat, t2_folded, re.I) for pat in continuation_patterns)
                    or any(kw in t2_folded for kw in EXTENDED_CONTINUATION_KEYWORDS)
                )
                if not is_cont:
                    lines2 = [l.strip() for l in t2.splitlines() if l.strip()]
                    cands2 = page_candidates(lines2, employees, cfg)
                    if not cands2 or cands2[0].score < 85:
                        is_cont = True

                if is_cont:
                    jobs.append(CertJob(job_id=len(jobs) + 1, src_pdf=pdf_path, pages=[0, 1], is_bundle=False))
                else:
                    jobs.append(CertJob(job_id=len(jobs) + 1, src_pdf=pdf_path, pages=[0], is_bundle=True))
                    jobs.append(CertJob(job_id=len(jobs) + 1, src_pdf=pdf_path, pages=[1], is_bundle=True))
            else:
                # Scanned 2-page document: almost always a 2-page certificate (front & back).
                # Skip slow synchronous OCR during preparation; Pass 1 handles extraction in parallel.
                jobs.append(CertJob(job_id=len(jobs) + 1, src_pdf=pdf_path, pages=[0, 1], is_bundle=False))
        else:
            # Multi-page bundle or multi-page certificate (>2 pages)
            infos = []
            for p_idx in range(n_pages):
                page = doc[p_idx]
                t = _get_page_text_cached(doc, p_idx, pdf_path, ocr_engine, text_cache)
                lines = [l.strip() for l in t.splitlines() if l.strip()]
                raw_cands = page_candidates(lines, employees, cfg)
                cands = cue_filter(raw_cands, lines, cfg)
                st, top, r = decide(cands, employees, cfg)
                is_start = bool(cands and cands[0].score >= accept_score)
                folded = fold(t).lower()
                cont = (
                    any(re.search(pat, folded, re.I) for pat in continuation_patterns)
                    or any(kw in folded for kw in EXTENDED_CONTINUATION_KEYWORDS)
                )
                if not is_start and not cont:
                    if not cands or cands[0].score < 85:
                        cont = True
                infos.append(
                    PageInfo(
                        decision=(st, top, r),
                        is_start=is_start,
                        number=None,
                        continuation=cont,
                        source="text" if page.get_text().strip() else "ocr",
                    )
                )
            raw_groups = group_pages(infos, employees, pages_per_cert, cfg)
            # If all pages form a single group covering the whole document, it is a single multi-page cert
            is_single_doc = (len(raw_groups) == 1 and len(raw_groups[0]["pages"]) == n_pages)
            for rg in raw_groups:
                zero_pages = [p - 1 for p in rg["pages"]]
                jobs.append(
                    CertJob(
                        job_id=len(jobs) + 1,
                        src_pdf=pdf_path,
                        pages=zero_pages,
                        is_bundle=not is_single_doc,
                    )
                )
        doc.close()
        if pdf_idx % 20 == 0:
            gc.collect()

    return jobs, roster_paths, roster_entries, aux_paths


def _extract_job_record(
    idx: int,
    job: CertJob,
    employees: list[Employee],
    cfg: dict[str, Any],
    ocr_engine: OcrEngine,
    roster_entries: list[dict[str, Any]],
    use_ai: bool,
    has_rescue_engine: bool,
    text_cache: dict[tuple[str, int], str],
) -> dict[str, Any]:
    doc = pymupdf.open(job.src_pdf)
    all_lines: list[str] = []
    first_p = job.pages[0] if job.pages else 0
    t0 = _get_page_text_cached(doc, first_p, job.src_pdf, ocr_engine, text_cache)
    for line in t0.splitlines():
        if line.strip():
            all_lines.append(line.strip())

    # Candidate matching on first page
    status, top, reason, clean_cands = find_bulk_candidates(
        all_lines, employees, cfg, stem_cue=job.src_pdf.stem
    )
    is_skp_off, lvl_off, peran_off = _detect_skp_offline("\n".join(all_lines))

    # Read remaining pages:
    # Always read if page has digital text (instant: 0.001s).
    # If scanned (no digital text), only OCR if page 0 was not accepted, or if SKP level/peran is missing.
    for p_idx in job.pages[1:]:
        p = doc[p_idx]
        has_text = bool(p.get_text().strip())
        if has_text or status != "accepted" or (is_skp_off and (not lvl_off or not peran_off)):
            t = _get_page_text_cached(doc, p_idx, job.src_pdf, ocr_engine, text_cache)
            for line in t.splitlines():
                if line.strip():
                    all_lines.append(line.strip())
            if status != "accepted":
                status, top, reason, clean_cands = find_bulk_candidates(
                    all_lines, employees, cfg, stem_cue=job.src_pdf.stem
                )
            if is_skp_off and (not lvl_off or not peran_off):
                _, lvl_sub, peran_sub = _detect_skp_offline("\n".join(all_lines))
                if lvl_sub:
                    lvl_off = lvl_sub
                if peran_sub:
                    peran_off = peran_sub

    doc.close()

    full_text = "\n".join(all_lines)
    detected_name = extract_raw_holder_name(all_lines, cfg) or ""

    # Certificate number extraction across issuers (SKP, Kemnaker, BNSP, Training)
    doc_num = ""
    m_c = re.search(r"(\d{4}\.[\d\.]+\.[A-Z0-9]+\.\d{2}\.\d{4})", full_text)
    if m_c:
        doc_num = m_c.group(1)
    else:
        m_kem = re.search(r"(?:No\.?\s*Reg(?:istrasi)?|Reg\.?|Registrasi)\s*[:]?\s*([0-9A-Z\/\.\-]{5,})", full_text, re.IGNORECASE)
        if m_kem:
            doc_num = m_kem.group(1).strip(" :-_.")
        else:
            m_bnsp = re.search(r"(?:MET|LSP)[\.\s]*\d{3}[\.\s]*\d{5,}\s*\d{4}", full_text, re.IGNORECASE)
            if m_bnsp:
                doc_num = m_bnsp.group(0).strip(" :-_.")
            else:
                m_c_gen = re.search(r"(?:Nomor\s+Sertifikat|Certificate\s+Number|No\.?\s*Sertifikat|No\.?\s*Certificate)[^\n\d]*(\d{3,}[\.\d\w\-\/]+)", full_text, re.IGNORECASE)
                if m_c_gen:
                    doc_num = m_c_gen.group(1).strip(" :-_.")

    # Title & Dates extraction
    doc_title = peran_off if is_skp_off else ""
    iss_date = ""
    exp_date = ""

    if not is_skp_off:
        m_prog = re.search(r"for successfully completed\s+(?:the\s+)?(.+?)(?:\s+and is able|\n|$)", full_text, re.IGNORECASE)
        if m_prog:
            clean_t = m_prog.group(1).replace("cours\xe9", "course").replace("coursé", "course")
            doc_title = clean_t.strip(" :-_.")
        else:
            tmpl_lines: list[str] = []
            t_val, _ = extract_title(all_lines, tmpl_lines, job.is_bundle, len(job.pages), cfg)
            if t_val and not any(bad in t_val.lower() for bad in ["the knowledge acquired", "is able to apply", "training in practice"]):
                doc_title = t_val
        iv, _, ev, _ = extract_dates(all_lines, cfg)
        if iv:
            iss_date = iv.isoformat()
        if ev:
            exp_date = ev.isoformat()

        # Relative validity period calculation (e.g. 'berlaku 3 tahun', 'valid for 3 years')
        if not exp_date and iss_date:
            m_rel = re.search(
                r"(?:berlaku|masa\s+berlaku|validity|valid\s+for)\s+(?:untuk\s+)?(?:jangka\s+waktu\s+)?([1-9])\s*(?:\([a-z]+\)\s*)?(?:tahun|years?)",
                full_text,
                re.IGNORECASE,
            )
            if m_rel:
                try:
                    years_add = int(m_rel.group(1))
                    iss_dt = datetime.fromisoformat(iss_date).date()
                    try:
                        exp_dt = iss_dt.replace(year=iss_dt.year + years_add)
                    except ValueError:
                        exp_dt = iss_dt.replace(year=iss_dt.year + years_add, day=28)
                    exp_date = exp_dt.isoformat()
                except Exception:
                    pass

    # Roster table enrichment (Dual-Key Matching: Cert No first, then Name)
    r_match = None
    if doc_num:
        r_match = match_roster("", roster_entries, cert_no=doc_num)
    if not r_match and top is not None:
        r_match = match_roster(employees[top.index].nama, roster_entries)
    if not r_match and detected_name:
        r_match = match_roster(detected_name, roster_entries)

    if r_match:
        is_skp_off = True
        if r_match.get("level"):
            lvl_off = r_match["level"]
        if r_match.get("cert_no") and not doc_num:
            doc_num = r_match["cert_no"]
        if r_match.get("peran") and not peran_off:
            peran_off = r_match["peran"]
            doc_title = peran_off

        # If holder matching against DB was unresolved, try matching roster tokens
        if top is None and r_match.get("tokens"):
            r_cands = page_candidates(r_match["tokens"], employees, cfg)
            if r_cands and r_cands[0].score >= 85.0:
                top = r_cands[0]
                status = "accepted"
                reason = ""

    # Dynamic text-based detection for transmittal letters / roster / auxiliary docs (handles generic filenames like SCAN_001)
    t_upper = full_text.upper()
    is_roster_doc_text = any(
        k in t_upper
        for k in [
            "LAMPIRAN SURAT PENGANTAR",
            "SURAT PENGANTAR PENGIRIMAN",
            "SURAT KETERANGAN PENGIRIMAN",
            "KODE OKUPASI JABATAN",
            "PENERBITANSERTIFIKASI",
            "PENERBITAN SERTIFIKASI",
            "DAFTAR PENERIMA",
            "DAFTAR SERTIFIKASI",
            "HASIL PERPANJANGAN",
            "HASIL PERPANJANGAN KEMNAKER",
        ]
    ) or ("KODE OKUPASI" in t_upper and ("NO SERTIFIKAT" in t_upper or "SERTIFIKAT" in t_upper)) or ("JENIS SERTIFIKASI" in t_upper and ("LISENSI" in t_upper or "SKP" in t_upper))

    aux_cat_detected = ""
    aux_text_map = [
        (["absen", "presensi", "daftar hadir", "attendance"], "Daftar Hadir"),
        (["invoice", "kwitansi", "receipt", "bukti bayar", "tagihan"], "Kwitansi / Invoice"),
        (["rundown", "jadwal", "agenda kegiatan", "susunan acara"], "Jadwal / Rundown"),
        (["surat tugas", "spt", "surat permohonan", "berita acara", "bast"], "Surat Tugas / Berita Acara"),
        (["evaluasi", "feedback", "kuesioner"], "Form Evaluasi"),
    ]
    t_lower = full_text.lower()
    if not is_roster_doc_text and "sertifikat" not in t_lower and "certificate" not in t_lower and "kompeten" not in t_lower:
        for keywords, cat in aux_text_map:
            if any(k in t_lower for k in keywords):
                aux_cat_detected = cat
                break

    if is_roster_doc_text or aux_cat_detected:
        status = "accepted"
        top = None
        clean_cands = []
        is_skp_off = False
        lvl_off = ""
        peran_off = ""
        doc_num = "-"
        iss_date = "-"
        exp_date = "-"
        if is_roster_doc_text:
            doc_title = "Surat Pengantar / Daftar Rekap"
            reason = "Dokumen rekap/surat pengantar otomatis terdeteksi dari isi teks"
        else:
            doc_title = f"Lampiran ({aux_cat_detected})"
            reason = f"Dokumen pendukung ({aux_cat_detected})"

    # Check if 100% resolved locally
    is_resolved_locally = (
        is_roster_doc_text
        or bool(aux_cat_detected)
        or (status == "accepted" and top is not None and (not is_skp_off or (lvl_off and peran_off)))
    )
    needs_ai = bool(use_ai and has_rescue_engine and not is_resolved_locally)

    return {
        "idx": idx,
        "job": job,
        "pdf_path": job.src_pdf,
        "all_lines": all_lines,
        "cands": clean_cands,
        "status": status,
        "top": top,
        "reason": reason,
        "detected_name": detected_name,
        "is_skp": is_skp_off,
        "level": lvl_off,
        "peran": peran_off,
        "doc_title": doc_title,
        "doc_num": doc_num,
        "iss_date": iss_date,
        "exp_date": exp_date,
        "needs_ai": needs_ai,
        "is_roster": is_roster_doc_text,
        "is_aux": bool(aux_cat_detected),
        "aux_cat": aux_cat_detected,
    }


def process_bulk_files(
    pdf_paths: list[Path],
    db_path: Path,
    out_dir: Path,
    use_ai: bool = True,
    gemini_api_key: str | None = None,
    gemini_model: str = "gemini-flash-lite-latest",
    run_name: str | None = None,
    progress_callback: Callable[[int, int, str], None] | None = None,
    use_ocr: bool = True,
) -> BulkResult:
    """Unified smart processor: handles 1 bundle, 20+ single files, or mixed folders with automatic roster enrichment."""
    cfg = load_config(None)
    employees = load_employees(db_path, cfg)
    ocr_engine = OcrEngine(dpi=130) if use_ocr else None

    if use_ai and not gemini_api_key:
        settings = load_user_settings()
        gemini_api_key = settings.get("gemini_api_key")
        if settings.get("model"):
            gemini_model = settings.get("model")

    rescue_engine = None
    if use_ai and gemini_api_key:
        rescue_engine = GeminiRescue(api_key=gemini_api_key, model=gemini_model)

    out_dir.mkdir(parents=True, exist_ok=True)
    if not run_name:
        timestamp_str = datetime.now().strftime("%d-%b-%Y_%H.%M")
        run_name = f"Hasil_Proses_{len(pdf_paths)}_Item_{timestamp_str}"

    target_run_dir = out_dir / run_name
    counter = 2
    while target_run_dir.exists():
        target_run_dir = out_dir / f"{run_name}_{counter}"
        counter += 1
    target_run_dir.mkdir(parents=True, exist_ok=True)

    pdf_out_dir = target_run_dir / "pdf"
    pdf_out_dir.mkdir(parents=True, exist_ok=True)

    if progress_callback:
        progress_callback(1, len(pdf_paths), "Menganalisis dokumen & mendeteksi struktur lampiran...")

    text_cache: dict[tuple[str, int], str] = {}

    # Step 1: Deconstruct into CertJobs & Roster docs
    jobs, roster_paths, roster_entries, aux_paths = prepare_cert_jobs(
        pdf_paths, employees, cfg, ocr_engine, text_cache=text_cache
    )
    total_jobs = len(jobs)

    # Copy roster documents to output
    for r_idx, r_path in enumerate(roster_paths, 1):
        target_name = "Lampiran_Surat_Pengantar.pdf" if len(roster_paths) == 1 else f"Lampiran_Surat_Pengantar_{r_path.stem}.pdf"
        shutil.copy2(r_path, pdf_out_dir / target_name)
        logger.info("Menyimpan dokumen roster ke: %s", target_name)

    # Copy auxiliary non-certificate documents to output
    aux_items: list[BulkItem] = []
    for a_idx, (a_path, a_cat) in enumerate(aux_paths, 1):
        clean_cat = clean_field(a_cat)
        target_aux_name = f"[LAMPIRAN]_{clean_cat}_{clean_field(a_path.stem)}.pdf"
        shutil.copy2(a_path, pdf_out_dir / target_aux_name)
        logger.info("Menyimpan lampiran non-sertifikat: %s -> %s", a_path.name, target_aux_name)
        aux_items.append(
            BulkItem(
                index=0,
                src_path=a_path,
                file_name=target_aux_name,
                nid="-",
                nama="-",
                jabatan="-",
                unit="-",
                is_skp=False,
                level="-",
                peran="-",
                judul=f"Lampiran ({a_cat})",
                number="-",
                issue_date="-",
                expiry_date="-",
                status="accepted",
                reason=f"Dokumen pendukung ({a_cat})",
                candidates=[],
                pages=[0],
                is_bundle=False,
            )
        )

    # =========================================================================
    # PASS 1: Extraction & Candidate Matching (Adaptive Multi-Worker)
    # =========================================================================
    workers = get_optimal_worker_count()
    completed_count = 0
    progress_lock = threading.Lock()

    def worker_fn(item: tuple[int, CertJob]) -> dict[str, Any]:
        nonlocal completed_count
        idx, job = item
        rec = _extract_job_record(
            idx=idx,
            job=job,
            employees=employees,
            cfg=cfg,
            ocr_engine=ocr_engine,
            roster_entries=roster_entries,
            use_ai=use_ai,
            has_rescue_engine=bool(rescue_engine),
            text_cache=text_cache,
        )
        with progress_lock:
            completed_count += 1
            if progress_callback:
                progress_callback(
                    completed_count,
                    total_jobs,
                    f"Membaca ({completed_count}/{total_jobs}): {job.src_pdf.name}",
                )
        return rec

    if total_jobs <= 2 or workers <= 1:
        pending_records = [worker_fn((idx, job)) for idx, job in enumerate(jobs, 1)]
    else:
        with concurrent.futures.ThreadPoolExecutor(max_workers=workers) as executor:
            records_unsorted = list(executor.map(worker_fn, list(enumerate(jobs, 1))))
        pending_records = sorted(records_unsorted, key=lambda r: r["idx"])

    gc.collect()

    # =========================================================================
    # PASS 2: Batch AI Rescue for Ambiguous/Unmatched Items (Max 5 per request)
    # =========================================================================
    items_for_ai = [r for r in pending_records if r["needs_ai"]]
    if items_for_ai and rescue_engine:
        batch_size = 5
        num_batches = (len(items_for_ai) + batch_size - 1) // batch_size
        for b_idx in range(num_batches):
            chunk = items_for_ai[b_idx * batch_size : (b_idx + 1) * batch_size]
            if progress_callback:
                progress_callback(
                    len(pending_records),
                    total_jobs,
                    f"AI Batch {b_idx + 1}/{num_batches} ({len(chunk)} sertifikat)...",
                )

            batch_docs: list[dict[str, Any]] = []
            open_docs: list[pymupdf.Document] = []
            try:
                for rec in chunk:
                    j: CertJob = rec["job"]
                    d = pymupdf.open(j.src_pdf)
                    open_docs.append(d)
                    pages_for_ai = [d[p_idx] for p_idx in j.pages]
                    batch_docs.append({
                        "id": rec["idx"],
                        "pages": pages_for_ai,
                        "candidates": rec["cands"],
                    })

                ai_batch_results = rescue_engine.batch_analyze_certificates(batch_docs, employees)
                for rec in chunk:
                    ai_res = ai_batch_results.get(rec["idx"]) or ai_batch_results.get(str(rec["idx"]))
                    if ai_res:
                        if ai_res.get("candidate") is not None:
                            rec["top"] = ai_res["candidate"]
                            rec["status"] = "accepted"
                            rec["reason"] = ""
                        if ai_res.get("detected_name"):
                            rec["detected_name"] = ai_res["detected_name"]
                            # Re-evaluate candidates against DB using the name detected by Gemini
                            ai_st, ai_top, ai_r, ai_cands = find_bulk_candidates(
                                [rec["detected_name"]], employees, cfg, stem_cue=rec["job"].src_pdf.stem
                            )
                            if ai_cands:
                                rec["cands"] = ai_cands
                            if ai_top is not None and ai_st == "accepted":
                                rec["top"] = ai_top
                                rec["status"] = "accepted"
                                rec["reason"] = ""
                            elif rec.get("top") is None:
                                rec["top"] = ai_top
                                rec["reason"] = ai_r
                        if ai_res.get("is_skp"):
                            rec["is_skp"] = True
                            if ai_res.get("level") and not rec["level"]:
                                rec["level"] = ai_res["level"]
                            if ai_res.get("peran") and not rec["peran"]:
                                rec["peran"] = ai_res["peran"]
                        if ai_res.get("title"):
                            rec["doc_title"] = ai_res["title"]
                        if ai_res.get("number") and not rec["doc_num"]:
                            rec["doc_num"] = ai_res["number"]
                        if ai_res.get("issue_date") and not rec["iss_date"]:
                            rec["iss_date"] = ai_res["issue_date"]
                        if ai_res.get("expiry_date") and not rec["exp_date"]:
                            rec["exp_date"] = ai_res["expiry_date"]
            finally:
                for d in open_docs:
                    try:
                        d.close()
                    except Exception:
                        pass
                gc.collect()

    # =========================================================================
    # PASS 3: Finalize Filenames, Save Output PDFs, and Write Reports
    # =========================================================================
    seen_names: dict[str, int] = {}
    items: list[BulkItem] = []

    for rec in pending_records:
        job: CertJob = rec["job"]
        pdf_path = rec["pdf_path"]
        status = rec["status"]
        top = rec["top"]
        reason = rec["reason"]
        detected_name = rec["detected_name"]
        is_skp = rec["is_skp"]
        level = rec["level"]
        peran = rec["peran"]
        doc_title = rec["doc_title"]
        cands = rec["cands"]

        emp_nid = ""
        emp_nama = ""
        emp_jabatan = ""
        emp_unit = ""

        is_roster = rec.get("is_roster", False)
        is_aux = rec.get("is_aux", False)
        aux_cat = rec.get("aux_cat", "")

        if is_roster:
            core_name = f"[LAMPIRAN]_Surat_Pengantar_{clean_field(pdf_path.stem)}"
            base_target = core_name
            emp_nama = "-"
            emp_nid = "-"
            emp_jabatan = "-"
            emp_unit = "-"
            doc_title = "Surat Pengantar / Daftar Rekap"
            status = "accepted"
            reason = "Dokumen rekap/surat pengantar otomatis terdeteksi dari isi teks"
        elif is_aux:
            clean_c = clean_field(aux_cat or "Dokumen_Pendukung")
            core_name = f"[LAMPIRAN]_{clean_c}_{clean_field(pdf_path.stem)}"
            base_target = core_name
            emp_nama = "-"
            emp_nid = "-"
            emp_jabatan = "-"
            emp_unit = "-"
            doc_title = f"Lampiran ({aux_cat})"
            status = "accepted"
            reason = f"Dokumen pendukung ({aux_cat})"
        elif status == "accepted" and top is not None:
            matched_emp = employees[top.index]
            emp_nid = matched_emp.nid
            emp_nama = matched_emp.nama
            emp_jabatan = matched_emp.jabatan
            emp_unit = matched_emp.unit
            status = "accepted"
            reason = ""
        else:
            # Validate detected_name so sentences like 'for successfully completed...' don't become person names
            cand_name = detected_name.strip() if detected_name else ""
            bad_name_terms = [
                "successfully", "completed", "course", "program", "certificate", "completion",
                "training", "knowledge", "acquired", "practice", "penerbitan", "kualifikasi",
                "presented to", "diberikan kepada", "menyatakan"
            ]
            if any(b in cand_name.lower() for b in bad_name_terms) or len(cand_name.split()) > 4:
                cand_name = ""

            emp_nama = cand_name or pdf_path.stem
            emp_jabatan = peran if is_skp else ""
            status = "unmatched"
            if not cands:
                reason = f"Nama '{emp_nama}' tidak ditemukan di database master"
            elif not reason:
                reason = "no confident holder"

        cand_dicts = []
        for c in cands[:5]:
            cand_dicts.append({
                "nid": employees[c.index].nid,
                "nama": employees[c.index].nama,
                "unit": employees[c.index].unit,
                "score": round(c.score, 1),
                "partial": c.partial,
            })

        if not doc_title:
            doc_title = peran if is_skp else "Sertifikat"

        # Build clean filename
        clean_nid_str = clean_field(emp_nid)
        clean_nama_str = clean_field(emp_nama)
        clean_level_str = clean_field(level)
        clean_peran_str = clean_field(peran or doc_title)
        clean_judul_str = clean_field(doc_title)

        if is_roster or is_aux:
            pass  # base_target already cleanly defined
        elif is_skp:
            parts = [p for p in [clean_nid_str, clean_nama_str, clean_level_str, clean_peran_str] if p]
            core_name = "_".join(parts) or clean_nama_str or "Sertifikat_SKP"
            base_target = core_name if status == "accepted" else f"[REVIEW]_{core_name}"
        else:
            parts = [p for p in [clean_nid_str, clean_nama_str, clean_judul_str] if p]
            core_name = "_".join(parts) or clean_nama_str or "Sertifikat"
            base_target = core_name if status == "accepted" else f"[REVIEW]_{core_name}"

        key = base_target.lower()
        count = seen_names.get(key, 0) + 1
        seen_names[key] = count
        final_file_name = f"{base_target}.pdf" if count == 1 else f"{base_target}_dup{count}.pdf"

        # Save output PDF
        target_pdf_path = pdf_out_dir / final_file_name
        if job.is_bundle:
            src_doc = pymupdf.open(job.src_pdf)
            new_doc = pymupdf.open()
            new_doc.insert_pdf(src_doc, from_page=job.pages[0], to_page=job.pages[-1])
            new_doc.save(target_pdf_path)
            new_doc.close()
            src_doc.close()
        else:
            shutil.copy2(job.src_pdf, target_pdf_path)

        item = BulkItem(
            index=rec["idx"],
            src_path=pdf_path,
            file_name=final_file_name,
            nid=emp_nid,
            nama=emp_nama,
            jabatan=emp_jabatan,
            unit=emp_unit,
            is_skp=is_skp,
            level=level,
            peran=peran,
            judul=doc_title,
            number=rec["doc_num"],
            issue_date=rec["iss_date"],
            expiry_date=rec["exp_date"],
            status=status,
            reason=reason,
            candidates=cand_dicts,
            pages=job.pages,
            is_bundle=job.is_bundle,
        )
        items.append(item)

    # Append auxiliary documents to items list
    for a_it in aux_items:
        a_it.index = len(items) + 1
        items.append(a_it)

    # 4. Write Single Combined data.xlsx and manifest_slice.xlsx
    _write_bulk_workbooks(target_run_dir, items)

    # 5. Write plan.json for ReviewDialog compatibility
    _write_bulk_plan_json(target_run_dir, items, db_path)

    accepted_cnt = sum(1 for it in items if it.status == "accepted")
    review_cnt = len(items) - accepted_cnt

    return BulkResult(
        run_dir=target_run_dir,
        total_files=len(items),
        accepted_count=accepted_cnt,
        review_count=review_cnt,
        items=items,
    )


def _write_bulk_workbooks(run_dir: Path, items: list[BulkItem]) -> None:
    """Generate a clean, professional unified data.xlsx and manifest_slice.xlsx."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rekap Sertifikat"

    headers = [
        "No",
        "File Asal",
        "Nama File Baru",
        "NID",
        "Nama Pegawai",
        "Jabatan",
        "Unit Kerja",
        "Tipe Dokumen",
        "Judul Sertifikat / Pelatihan",
        "Level (Khusus SKP)",
        "Peran Kompetensi (Khusus SKP)",
        "Nomor Sertifikat",
        "Tanggal Terbit",
        "Tanggal Expired",
        "Status Validasi",
        "Keterangan / Alasan Review",
    ]

    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    border_thin = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = border_thin

    row_fill_even = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    row_fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    status_font_ok = Font(name="Segoe UI", size=9, bold=True, color="047857")
    status_font_rev = Font(name="Segoe UI", size=9, bold=True, color="B45309")

    for row_idx, it in enumerate(items, 2):
        fill = row_fill_even if row_idx % 2 == 0 else row_fill_odd
        vals = [
            it.index,
            it.src_path.name,
            it.file_name,
            it.nid,
            it.nama,
            it.jabatan,
            it.unit,
            "SKP Ketenagalistrikan" if it.is_skp else "Sertifikat Standar",
            it.judul,
            it.level,
            it.peran,
            it.number,
            it.issue_date,
            it.expiry_date,
            "Tervalidasi" if it.status == "accepted" else "Perlu Review",
            it.reason if it.status != "accepted" else "-",
        ]
        for col_idx, v in enumerate(vals, 1):
            c = ws.cell(row=row_idx, column=col_idx, value="" if v is None else str(v))
            c.font = Font(name="Segoe UI", size=9)
            c.fill = fill
            c.border = border_thin
            if col_idx in [1, 4, 13, 14]:
                c.alignment = align_center
            else:
                c.alignment = align_left
            if col_idx == 15:
                c.font = status_font_ok if it.status == "accepted" else status_font_rev
                c.alignment = align_center

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 45)

    ws.row_dimensions[1].height = 26

    wb.save(run_dir / "data.xlsx")
    wb.save(run_dir / "manifest_slice.xlsx")
    wb.close()


def _write_bulk_plan_json(run_dir: Path, items: list[BulkItem], db_path: Path) -> None:
    """Save plan.json so ReviewDialog can view and correct any unmatched items."""
    plan_groups = []
    for it in items:
        plan_groups.append({
            "id": it.index,
            "pages": [p + 1 for p in it.pages],
            "status": it.status,
            "reason": it.reason,
            "nid": it.nid,
            "nama": it.nama,
            "raw_name": it.nama,
            "jabatan": it.jabatan,
            "unit": it.unit,
            "candidates": it.candidates,
            "fields": {
                "title": {"value": it.judul, "confidence": "high"},
                "issue_date": {"value": it.issue_date, "confidence": "high"},
                "expiry_date": {"value": it.expiry_date, "confidence": "high"},
                "number": {"value": it.number, "confidence": "high"},
                "is_skp": {"value": it.is_skp, "confidence": "high"},
                "level": {"value": it.level, "confidence": "high"},
                "peran": {"value": it.peran, "confidence": "high"},
            },
            "file_name": it.file_name,
            "review_file_name": it.file_name if it.status != "accepted" else None,
            "src_pdf": str(it.src_path),
            "is_bundle": it.is_bundle,
        })

    plan_data = {
        "version": 1,
        "mode": "bulk",
        "input_pdf": str(items[0].src_path) if items else "",
        "database": str(db_path),
        "page_count": len(items),
        "groups": plan_groups,
    }
    (run_dir / "plan.json").write_text(json.dumps(plan_data, ensure_ascii=False, indent=2), encoding="utf-8")


def update_bulk_workbooks_from_groups(run_dir: Path, groups: list[dict[str, Any]]) -> None:
    """Rebuild data.xlsx and manifest_slice.xlsx from plan.json groups after user approval in review dialog."""
    items: list[BulkItem] = []
    for g in groups:
        f = g.get("fields", {})
        is_skp_val = f.get("is_skp", {}).get("value")
        items.append(
            BulkItem(
                index=g.get("id", 1),
                src_path=Path(g.get("src_pdf", "")),
                file_name=g.get("file_name", ""),
                nid=g.get("nid") or "",
                nama=g.get("nama") or "",
                jabatan=g.get("jabatan") or "",
                unit=g.get("unit") or "",
                is_skp=bool(is_skp_val),
                level=f.get("level", {}).get("value") or "",
                peran=f.get("peran", {}).get("value") or "",
                judul=f.get("title", {}).get("value") or "",
                number=f.get("number", {}).get("value") or "",
                issue_date=f.get("issue_date", {}).get("value") or "",
                expiry_date=f.get("expiry_date", {}).get("value") or "",
                status=g.get("status", "accepted"),
                reason=g.get("reason", ""),
                candidates=g.get("candidates", []),
                pages=[p - 1 for p in g.get("pages", [1])],
                is_bundle=bool(g.get("is_bundle", False)),
            )
        )
    _write_bulk_workbooks(run_dir, items)

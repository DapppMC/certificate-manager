from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Any

import openpyxl

MANIFEST_HEADERS = [
    "No",
    "NID",
    "Nama Pegawai",
    "Judul Sertifikasi/Pelatihan",
    "Jenis Dokumen (Sertifikasi/Pelatihan)",
    "Lembaga/Penyelenggara",
    "Nomor Sertifikat",
    "Tanggal Terbit",
    "Tanggal Expired",
    "Nama File PDF",
]

UNMATCHED_HEADERS = [
    "Halaman",
    "Alasan",
    "Kandidat 1 Nama",
    "Kandidat 1 NID",
    "Kandidat 1 Unit",
    "Kandidat 1 Skor",
    "Kandidat 2 Nama",
    "Kandidat 2 NID",
    "Kandidat 2 Unit",
    "Kandidat 2 Skor",
    "Kandidat 3 Nama",
    "Kandidat 3 NID",
    "Kandidat 3 Unit",
    "Kandidat 3 Skor",
    "Judul",
    "Nomor Sertifikat",
    "Tanggal Terbit",
    "Tanggal Expired",
    "Sumber Teks",
]

NEEDS_REVIEW_HEADERS = [
    "Nama File PDF",
    "NID",
    "Halaman",
    "Kolom",
    "Nilai",
    "Keyakinan",
]

FIELD_TO_HEADER = {
    "title": "Judul Sertifikasi/Pelatihan",
    "doc_type": "Jenis Dokumen (Sertifikasi/Pelatihan)",
    "issuer": "Lembaga/Penyelenggara",
    "number": "Nomor Sertifikat",
    "issue_date": "Tanggal Terbit",
    "expiry_date": "Tanggal Expired",
}


def text_cell(ws: openpyxl.worksheet.worksheet.Worksheet, row: int, col: int, value: Any) -> None:
    cell = ws.cell(row=row, column=col, value="" if value is None else str(value))
    cell.number_format = "@"


def _format_date(iso_str: str | None, date_fmt: str) -> str:
    if not iso_str:
        return ""
    try:
        d = dt.date.fromisoformat(iso_str)
        return d.strftime(date_fmt)
    except Exception:
        return iso_str


def _page_label(pages: list[int]) -> str:
    if not pages:
        return ""
    if len(pages) == 1:
        return str(pages[0])
    return f"{pages[0]}-{pages[-1]}"


def write_workbooks(run_dir: Path, groups: list[dict[str, Any]], cfg: dict[str, Any]) -> None:
    date_fmt = cfg.get("output", {}).get("date_format", "%d/%m/%Y")

    # 1. manifest_slice.xlsx
    wb_m = openpyxl.Workbook()
    ws_m = wb_m.active
    ws_m.title = "manifest"

    for col_idx, h in enumerate(MANIFEST_HEADERS, 1):
        text_cell(ws_m, 1, col_idx, h)

    accepted_groups = [g for g in groups if g["status"] == "accepted"]
    for row_idx, grp in enumerate(accepted_groups, 2):
        no_val = row_idx - 1
        ws_m.cell(row=row_idx, column=1, value=no_val)
        text_cell(ws_m, row_idx, 2, grp.get("nid", ""))
        text_cell(ws_m, row_idx, 3, grp.get("nama", ""))
        text_cell(ws_m, row_idx, 4, grp["fields"]["title"]["value"] or "")
        text_cell(ws_m, row_idx, 5, grp["fields"]["doc_type"]["value"] or "")
        text_cell(ws_m, row_idx, 6, grp["fields"]["issuer"]["value"] or "")
        text_cell(ws_m, row_idx, 7, grp["fields"]["number"]["value"] or "")
        text_cell(ws_m, row_idx, 8, _format_date(grp["fields"]["issue_date"]["value"], date_fmt))
        text_cell(ws_m, row_idx, 9, _format_date(grp["fields"]["expiry_date"]["value"], date_fmt))
        text_cell(ws_m, row_idx, 10, grp.get("file_name", ""))

    wb_m.save(run_dir / "manifest_slice.xlsx")
    wb_m.close()

    # 2. unmatched_report.xlsx
    wb_u = openpyxl.Workbook()
    ws_u = wb_u.active
    ws_u.title = "unmatched"
    for col_idx, h in enumerate(UNMATCHED_HEADERS, 1):
        text_cell(ws_u, 1, col_idx, h)

    unmatched_groups = [g for g in groups if g["status"] == "unmatched"]
    for row_idx, grp in enumerate(unmatched_groups, 2):
        cands = grp.get("candidates", [])
        text_cell(ws_u, row_idx, 1, _page_label(grp.get("pages", [])))
        text_cell(ws_u, row_idx, 2, grp.get("reason", ""))
        for k in range(3):
            base_col = 3 + k * 4
            if k < len(cands):
                c = cands[k]
                text_cell(ws_u, row_idx, base_col, c.get("nama", ""))
                text_cell(ws_u, row_idx, base_col + 1, c.get("nid", ""))
                text_cell(ws_u, row_idx, base_col + 2, c.get("unit", ""))
                score = c.get("score")
                text_cell(ws_u, row_idx, base_col + 3, str(score) if score is not None else "")
            else:
                text_cell(ws_u, row_idx, base_col, "")
                text_cell(ws_u, row_idx, base_col + 1, "")
                text_cell(ws_u, row_idx, base_col + 2, "")
                text_cell(ws_u, row_idx, base_col + 3, "")
        text_cell(ws_u, row_idx, 15, grp["fields"]["title"]["value"] or "")
        text_cell(ws_u, row_idx, 16, grp["fields"]["number"]["value"] or "")
        text_cell(ws_u, row_idx, 17, _format_date(grp["fields"]["issue_date"]["value"], date_fmt))
        text_cell(ws_u, row_idx, 18, _format_date(grp["fields"]["expiry_date"]["value"], date_fmt))
        text_cell(ws_u, row_idx, 19, grp.get("source", ""))

    ws_nr = wb_u.create_sheet(title="needs_review")
    for col_idx, h in enumerate(NEEDS_REVIEW_HEADERS, 1):
        text_cell(ws_nr, 1, col_idx, h)

    nr_row_idx = 2
    for grp in accepted_groups:
        file_name = grp.get("file_name", "")
        nid = grp.get("nid", "")
        halaman = _page_label(grp.get("pages", []))
        fields = grp.get("fields", {})
        for field_key in ["title", "doc_type", "issuer", "number", "issue_date", "expiry_date"]:
            f_info = fields.get(field_key, {})
            if f_info.get("confidence") == "low":
                header_name = FIELD_TO_HEADER.get(field_key, field_key)
                raw_val = f_info.get("value")
                if field_key in ("issue_date", "expiry_date"):
                    formatted_val = _format_date(raw_val, date_fmt)
                else:
                    formatted_val = "" if raw_val is None else str(raw_val)
                text_cell(ws_nr, nr_row_idx, 1, file_name)
                text_cell(ws_nr, nr_row_idx, 2, nid)
                text_cell(ws_nr, nr_row_idx, 3, halaman)
                text_cell(ws_nr, nr_row_idx, 4, header_name)
                text_cell(ws_nr, nr_row_idx, 5, formatted_val)
                text_cell(ws_nr, nr_row_idx, 6, "low")
                nr_row_idx += 1

    wb_u.save(run_dir / "unmatched_report.xlsx")
    wb_u.close()
